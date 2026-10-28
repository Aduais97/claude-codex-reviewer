#!/usr/bin/env python3
"""HNZ RFP 0018709 — draft answers Part 2: Capability Q2-Q6, Capacity P1-P6, H&S H1-H5.

Companion to build_draft_answers.py (which covers EB Q1-Q3 + Capability Q1/Q7).
[PLACEHOLDERS] highlighted yellow; red drafting notes to delete before submission.
Dash-free body text per ASG client-doc standard.

Usage: .venv/bin/python build_draft_answers_part2.py <out.docx>
"""
import re, sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_COLOR_INDEX

NAVY = RGBColor(0x1B, 0x2A, 0x4A)

def add_runs(p, text, size=10.5):
    for part in re.split(r"(\[[^\]]+\])", text):
        if not part:
            continue
        run = p.add_run(part)
        run.font.size = Pt(size)
        run.font.name = "Calibri"
        if part.startswith("["):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            run.bold = True

def para(doc, text, size=10.5, bold=False, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if bold:
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(size)
    else:
        add_runs(p, text, size)
    return p

def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    add_runs(p, text)

def heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = NAVY
    return h

def note(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("DRAFTING NOTE (delete before submission): " + text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)

doc = Document()

heading(doc, "HNZ RFP 0018709 - Draft Answers v1, Part 2", 0)
para(doc, "Capability Q2-Q6, Capacity Q1-Q6 and Health & Safety Q1-Q5, for Response Form V3 blue sections. Respondent: Absolute Security Group Limited. Districts: Auckland, Counties Manukau, Waitemata, Northland. Part 1 covers Economic Benefits Q1-Q3 and Capability Q1/Q7. Yellow [PLACEHOLDERS] need Ahmad's input; red notes are internal and must be deleted.", 10)

# ================= CAPABILITY Q2 =================
heading(doc, "CAPABILITY - Q2: Screening process for healthcare security officer candidates, including physical fitness and English language standards", 1)
para(doc, "Every candidate passes a single documented pipeline before they are eligible for any client site. No step is waived for urgency. For Health NZ assignments three additional gates apply at the end of the pipeline.", space_after=4)
para(doc, "Standard ASG screening pipeline:", bold=True, space_after=2)
bullet(doc, "Identity and right-to-work verification, including original document sighting.")
bullet(doc, "Licence verification: current full Certificate of Approval, category Property Guard, verified against the PSPLA public register, with NZQA unit standards US 27364, US 27360 and US 27361 confirmed complete. Candidates holding only Crowd Controller class are not deployed to guarding roles until the Property Guard class is held.")
bullet(doc, "Ministry of Justice criminal record check at ASG's cost; adverse results assessed against a documented relevance framework.")
bullet(doc, "Two reference checks from previous employers or supervisors, completed by phone and recorded.")
bullet(doc, "Structured interview including scenario-based judgement questions on de-escalation, use of force thresholds and customer service under pressure.")
bullet(doc, "Physical fitness and ability: [DESCRIBE actual method, e.g. a practical functional assessment covering sustained standing/walking, stair work, response pace, and controlled physical intervention technique during training; or a health declaration plus practical assessment at induction]. The standard is the ability to perform all duties of a full shift including physical intervention where required.")
bullet(doc, "English language: spoken and written English assessed against the Immigration New Zealand Skilled Migrant Category standard. Evidence is held per officer: [IELTS or equivalent test result / NZ or comparable English-medium qualification / documented ASG written and verbal assessment]. Officers whose evidence is incomplete are supported to obtain it before assignment, not screened out by default.")
bullet(doc, "Presentation standard: clean and approachable appearance is a condition of deployment, assessed at interview and enforced at every shift. We note and support Health NZ's direction away from militaristic, all-black or dark-blue presentation and will align uniform and grooming standards with each District Schedule.")
para(doc, "Additional gates for Health NZ assignment:", bold=True, space_after=2)
bullet(doc, "Occupational infectious disease screening and vaccination to Category A requirements per Attachments W and X, completed before first shift, with an evidence register maintained under the Privacy Act 2020 and Health Information Privacy Code 2020 and available to Health NZ for audit or contact tracing.")
bullet(doc, "Children's Act 2014 safety checks (identity, Police vet, referee checks, interview, risk assessment) completed before the Start Date for any officer whose role meets the Children's Worker definition, with three-yearly rechecks.")
bullet(doc, "Health NZ acceptance: we support per-officer introductory meetings where a District requires them, and will replace any officer Health NZ declines, at no additional cost, within the agreed timeframe.")

# ================= CAPABILITY Q3 =================
heading(doc, "CAPABILITY - Q3: Induction, onboarding and site-specific training processes", 1)
note(doc, "Question says these should be ASG's processes, not the sites'. Keep HNZ inductions as the layered final step only.")
para(doc, "ASG onboarding (every officer, before any solo deployment):", bold=True, space_after=2)
bullet(doc, "Day one: employment agreement on ASG terms, code of conduct, uniform and equipment issue, payroll and rostering setup (Deputy), privacy and confidentiality undertaking.")
bullet(doc, "Core induction: ASG field training manual covering legal powers and limits (trespass, defence of self and others, citizen's arrest thresholds), de-escalation doctrine, incident reporting standards, communications protocol, customer service expectations. [X hours over X days].")
bullet(doc, "Incident platform training: every officer is trained and assessed on writing evidence-quality reports in our digital incident system before first shift.")
bullet(doc, "Supervised buddy shifts: [N] shifts alongside an experienced officer with a written competency sign-off by a supervisor before solo deployment.")
para(doc, "Site-specific induction (every officer, every new site):", bold=True, space_after=2)
bullet(doc, "A site induction pack per site: layout and access points, site SOPs and post orders, escalation contacts, known risks, reporting expectations. Acknowledged in writing by the officer.")
bullet(doc, "First shifts at a new site are supervised or paired until the supervisor signs off site competency against a checklist.")
bullet(doc, "Site SOPs are reviewed with the client at least annually and within one month of any identified gap, consistent with the District Schedule requirements.")
para(doc, "For Health NZ sites this ASG framework is layered beneath the required Health NZ inductions before deployment: the Health NZ contractor site induction, security induction (operating, escalation and emergency procedures, site familiarisation) and contractor health and safety induction. Where a District requires system competency (CCTV monitoring, Gallagher access control, radio protocols per district security SOPs) this is completed and verified before the officer is counted in the deployable pool. Officers assigned to Mental Health facilities complete SPEC (full four-day course) before permanent assignment, with annual revalidation, and we notify the District of any role change or departure so replacement training is arranged ahead of need.", space_after=4)

# ================= CAPABILITY Q4 =================
heading(doc, "CAPABILITY - Q4: Ongoing training and development, frequency, and pipeline into senior roles", 1)
para(doc, "Ongoing training (all officers):", bold=True, space_after=2)
bullet(doc, "Scenario-based de-escalation refreshers [FREQUENCY, e.g. quarterly toolbox sessions plus an annual full-day refresher], built around real incidents from our own reporting data.")
bullet(doc, "Legal refreshers on powers, use of force and trespass [FREQUENCY].")
bullet(doc, "Report writing and evidence quality coaching, driven by supervisor review of live incident reports.")
bullet(doc, "First aid: [CURRENT STATUS, e.g. X% of officers hold a current First Aid certificate; renewals tracked in the training register].")
bullet(doc, "Specialist streams: loss prevention and covert tradecraft, control room and CCTV operation, event command. Officers apply into these streams after [X months] of strong performance.")
bullet(doc, "Training investment: [X hours / $X] per officer per year beyond licensing minimums, tracked per officer in a training register with expiry alerts, producible to Health NZ on request.")
para(doc, "Under this contract, annual revalidation requirements (HSO Violence Reduction Training at the level set by each District Security Manager, and SPEC for Mental Health assigned officers) are scheduled through the same register, with backfill planned ahead of course dates so coverage is not interrupted.", space_after=4)
para(doc, "Pipeline into senior roles:", bold=True, space_after=2)
bullet(doc, "The pathway is explicit: frontline officer, specialist capability, senior officer, supervisor, operations management. Progression criteria are published internally so officers know what the next step requires.")
bullet(doc, "[N] of our current senior and supervisory staff began with ASG as frontline guards, including [EXAMPLE: role, years, anonymised or named].")
bullet(doc, "For this contract the pathway extends into healthcare specialisation: officers who complete HSO Violence Reduction Training and accrue hospital experience become eligible for senior healthcare roles and the Northland Security Operations Centre stream, which gives Health NZ a deepening bench rather than a static one.")

# ================= CAPABILITY Q5 =================
heading(doc, "CAPABILITY - Q5: Processes for ensuring compliance with legislation, regulatory requirements, industry standards, licences, certifications, Health NZ policies and site requirements", 1)
para(doc, "Compliance at ASG is systematised, not left to memory:", space_after=4)
bullet(doc, "Licence and certification register: every officer's COA class, expiry, NZQA units, first aid, SPEC/SI status and English evidence is held in a central register with automated expiry alerts at [60/30/7] days. Rostering is blocked for any officer whose COA has lapsed; there is no manual override below director level.")
bullet(doc, "Company licensing: ASG holds a current PSPLA Company Licence (class Guarding) [LICENCE NUMBER]; a copy is available to Health NZ on request at any time.")
bullet(doc, "Employment law: payroll runs through PayHero with Holidays Act compliant leave and public holiday calculations; employment agreements, minimum wage movements and KiwiSaver obligations are reviewed [FREQUENCY/ADVISOR]. We acknowledge the contract obligation to report any Labour Inspectorate investigation or adverse determination, and we operate so that there is nothing to report.")
bullet(doc, "Privacy: officers are trained on the Privacy Act 2020 at induction; for Health NZ sites this extends to the Health Information Privacy Code 2020, with the higher standard of care for health information built into incident reporting training (no patient identifiers beyond what the report lawfully requires, secure platform access, NZ-hosted data [CONFIRM hosting status]).")
bullet(doc, "Health NZ policies and site requirements: policies notified to us are converted into SOP and induction content with officer acknowledgement workflows, so compliance is auditable per officer. Site rules (radio protocols, equipment return at end of shift, ID display, phone restrictions) are enforced by supervisors and spot checks.")
bullet(doc, "Weekly service reporting: our platform produces weekly reports of services delivered, hours, sites and incidents per district, meeting the RFP's weekly reporting requirement and each District Schedule's stated cadence.")
bullet(doc, "Assurance: [INTERNAL AUDIT PRACTICE, e.g. monthly supervisor site inspections, quarterly compliance self-audit, annual external review; state what is real]. Findings are tracked to closure with named owners.")

# ================= CAPABILITY Q6 =================
heading(doc, "CAPABILITY - Q6: Quality assurance for subcontractors", 1)
para(doc, "ASG does not use subcontractors to deliver security officer services. Every officer deployed under this contract will be directly employed by Absolute Security Group Limited on an ASG employment agreement. We take this position for the same reasons Health NZ stated in the supplier briefing: accountability and control, consistency of service delivery, health and safety, training and competency, vetting and compliance, workforce continuity, and employment standards.", space_after=4)
para(doc, "We acknowledge and accept that under this contract subcontractors may only be used during a force majeure event and with the express approval of the District Security Manager, that this rule is universal across all Districts, and that no exceptions apply. In such an event, any temporary resource would be engaged only with that express approval, briefed against the relevant site SOPs, supervised by ASG, and subject to the same licensing verification before deployment. Outside that narrow circumstance, our delivery model is one hundred percent directly employed staff.", space_after=4)

# ================= CAPACITY P1 =================
heading(doc, "CAPACITY - Q1: Current staffing capacity per district, healthcare-relevant staff, and current clients per district", 1)
note(doc, "HARD RULE from the review: only count officers verified in the COA-class audit (Property Guard, not Crowd Control) with NZQA units and English evidence, NET of hours committed to existing contracts in the same time bands. Every number here is testable at reference checks and per-officer vetting meetings.")
para(doc, "ASG employs [N] active security staff across the Auckland region [GROUP OR ENTITY FIGURE, confirm entity attribution]. For this contract the relevant figure is deployable capacity: officers verified as holding a full Property Guard COA with NZQA US 27364, US 27360 and US 27361 and documented English evidence, net of existing commitments.", space_after=4)
para(doc, "Current verified deployable capacity by district:", bold=True, space_after=2)
bullet(doc, "Auckland District: [N] officers living within urgent response range of Auckland City Hospital and Greenlane, of whom [N] have experience relevant to healthcare security (mental distress response, intoxication management, physical intervention under a de-escalation first doctrine).")
bullet(doc, "Counties Manukau District: [N] officers resident in South Auckland, of whom [N] have relevant experience as above.")
bullet(doc, "Waitemata District: [N] officers resident in West Auckland and the North Shore, of whom [N] have relevant experience.")
bullet(doc, "Northland District: ASG does not currently maintain a Northland workforce. Our Northland capacity is a dated local recruitment and training plan (set out in Q6 and in our mobilisation plan) building a Te Tai Tokerau based team of [N] officers plus Security Operations Centre staffing, fully deployable between 1 April and 1 July 2027. We are deliberate about this honesty: Health NZ confirmed at the supplier briefing that readiness at the point of deployment, with a clear plan, is what will be assessed.")
para(doc, "In addition to current verified staff, we will stand up a dedicated Health NZ pool of [N] officers by contract commencement, recruited through the community pipelines described in our Economic Benefits response, so that Health NZ work is served by a committed pool rather than the margins of our existing roster.", space_after=4)
para(doc, "Current clients per district [MARK THIS TABLE COMMERCIALLY SENSITIVE]:", bold=True, space_after=2)
bullet(doc, "Auckland District: [CLIENT LIST with service type and years, e.g. city-centre precinct security programme since (YEAR); entertainment and hospitality venues; corporate sites. Confirm naming decisions.]")
bullet(doc, "Counties Manukau District: [CLIENT LIST or honest statement of current footprint.]")
bullet(doc, "Waitemata District: [CLIENT LIST or honest statement.]")
bullet(doc, "Northland District: none at present, consistent with the recruitment-led plan above.")

# ================= CAPACITY P2 =================
heading(doc, "CAPACITY - Q2: Ongoing recruitment strategy per district and community-reflective staffing", 1)
para(doc, "Our recruitment model is continuous, not reactive: we maintain a standing pipeline so that growth and attrition are absorbed without service impact.", space_after=4)
bullet(doc, "Channels: community and whanau networks within South and West Auckland (our highest-yield channel), [MSD/Work and Income relationships if real], industry transfers attracted by our training pathway, and the iwi and Pacific Hauora partnerships described in Economic Benefits Q1.")
bullet(doc, "Pipeline mechanics: applications are processed weekly; a candidate with an existing Property Guard COA is deployable in [X] weeks (screening, health screening and inductions); a candidate requiring a new COA in [X] weeks. For Health NZ we will run recruitment waves sized to the mobilisation plan, with Category A health screening started at preferred-supplier stage because it takes two to six weeks per officer.")
bullet(doc, "Auckland District: recruitment targets officers resident within urgent-response range of the central hospitals, weighted to maintain night and weekend availability.")
bullet(doc, "Counties Manukau District: recruitment through South Auckland community networks so the Middlemore-facing pool reflects the Pasifika and Maori communities it serves; language capability ([LANGUAGES]) is treated as a selection asset.")
bullet(doc, "Waitemata District: West Auckland and North Shore resident hiring for response coverage at both hospital campuses.")
bullet(doc, "Northland District: local-first hiring in Whangarei, Kaitaia, Kawakawa and Dargaville through [IWI/TRAINING PARTNER, ONLY IF CONTACTED, else 'a Te Tai Tokerau training partnership to be formalised by (DATE)'], with COA training support for strong local candidates who are not yet licensed. We will not staff Northland by rotating Auckland officers through remote posts; that model fails on fatigue, cost and community trust.")
bullet(doc, "Reflectiveness is measured, not asserted: pool composition versus district demographics is reported at QBRs (see Economic Benefits Q2).")

# ================= CAPACITY P3 =================
heading(doc, "CAPACITY - Q3: Backfill and surge processes, points of contact, operating hours", 1)
para(doc, "ASG operates 24 hours a day, 365 days a year. [DESCRIBE REAL STRUCTURE: staffed operations desk hours + after-hours duty manager, or 24/7 staffed desk]. There is one phone number and one email for Health NZ requests at any hour, answered by a person, not an automated service, meeting the Northland live-answer requirement.", space_after=4)
para(doc, "Points of contact and escalation:", bold=True, space_after=2)
bullet(doc, "First contact: ASG Operations [PHONE / EMAIL], 24/7.")
bullet(doc, "Escalation: duty Operations Manager [NAME/ROLE], then [OPS MANAGER NAME], then Ahmad Duais, CEO. Escalation to director level is available to Health NZ at any hour.")
para(doc, "Backfill process (scheduled posts):", bold=True, space_after=2)
bullet(doc, "Absence identified (sick call or no-show alert in Deputy) triggers an immediate offer cascade to the qualified pool for that district; acceptance is confirmed and the District Security Manager notified with the replacement officer's name and ETA. Target: replacement confirmed within [X] minutes of the absence being known.")
bullet(doc, "Every scheduled post has a named pool of cross-inducted officers (the contract's Pool concept) sized to cover expected leave and absence, so backfill comes from officers who already hold the site induction.")
para(doc, "Surge and on-demand response:", bold=True, space_after=2)
bullet(doc, "Tier 1 (one to two officers): dispatched from on-shift flex capacity or the on-call list. Performance against the contract KPI (90% within two hours, 100% within four) is tracked per event in our platform.")
bullet(doc, "Tier 2 (three or more officers): duty manager activates the recall list and, where appropriate, draws from adjacent districts. Evidence of this capability: [REAL EXAMPLE with numbers, e.g. stood up N officers in X hours for (event), confirm disclosure].")
bullet(doc, "District response commitments: Auckland urgent requests within 120 minutes; Counties Manukau within the timeframe set per request (one to twenty-four hours); Waitemata per request; Northland within 45 minutes via town-based on-call pools in each of the four hospital towns, which is a function of the local-hiring model described in Q2.")
bullet(doc, "Honest prioritisation: where simultaneous demands occur across clients, dispatch follows contractual response commitments first, then severity. The dedicated Health NZ pool exists precisely so that Health NZ response does not compete with our existing clients' peaks.")

# ================= CAPACITY P4 =================
heading(doc, "CAPACITY - Q4: Staff turnover rates per district and improvement actions", 1)
note(doc, "Numbers must be reproducible from payroll records. If district-level splits are not meaningful at current scale, say so and give the company figure, this reads better than invented splits.")
para(doc, "Our annualised staff turnover for the twelve months to [DATE] was [X]% ([METHOD: leavers divided by average headcount]). At our current scale we measure this at company level rather than by district; we commit to reporting Health NZ pool turnover by district at each QBR from commencement. For context, reported turnover in the New Zealand security industry commonly runs well above [X]% [OPTIONAL: cite NZSA or omit].", space_after=4)
para(doc, "We do not yet operate healthcare sites; our commitment is transparency: Health NZ will see pool-specific turnover from the first quarter, not a blended company number.", space_after=4)
para(doc, "What we do to keep people:", bold=True, space_after=2)
bullet(doc, "Pay and conditions: [REAL POSITION, e.g. rates above licensing-industry entry level, guaranteed minimum hours for pool members, penal rates passed through].")
bullet(doc, "A visible career pathway (Capability Q4): people stay where there is somewhere to go.")
bullet(doc, "Training investment and specialist streams, which for this contract includes Health NZ funded HSO Violence Reduction Training, a genuine retention asset.")
bullet(doc, "Engagement: [REAL PRACTICES: regular check-ins, EAP access, recognition, roster fairness. Delete what is not true.]")
bullet(doc, "Measured effect: [TREND if available, e.g. turnover down from X% to Y% over two years; otherwise delete].")
para(doc, "We are also direct about the contract mechanics: because Health NZ funds initial training, we carry the replacement training cost for any officer who leaves within three months. Our onboarding quality and pool design are built to make that a rare event, and we accept the incentive alignment it creates.", space_after=4)

# ================= CAPACITY P5 =================
heading(doc, "CAPACITY - Q5: Account management structure with national and district contacts, and half-page bios", 1)
para(doc, "Structure (single national owner, named district contacts, director access at all times):", space_after=4)
bullet(doc, "National relationship owner: Ahmad Duais, Chief Executive. Attends Quarterly Business Reviews with the Strategic Category Manager, owns KPI reporting including Pool status, and is the escalation end-point.")
bullet(doc, "Operations Manager, Auckland metro districts: [NAME]. Day-to-day contact for the Auckland, Counties Manukau and Waitemata District Security Managers; attends the Counties Manukau monthly meetings and Waitemata fortnightly meetings; owns rosters, backfill performance and incident follow-up.")
bullet(doc, "Northland District Manager: [NAME / recruitment status, target appointment by (DATE)]. Based in Whangarei; attends fortnightly local meetings and, with the CEO, quarterly senior meetings per the Northland schedule; owns the four-town pools and the SOC relationship.")
bullet(doc, "Security Operations Centre lead (Northland): [NAME / status]; controller roster, SOC standards and system competency.")
bullet(doc, "Healthcare security specialist: [NAME / status]; clinical-environment practice, patient watch SOPs, training alignment with the HSO Violence Reduction programme.")
para(doc, "Bios (half page each):", bold=True, space_after=2)
bullet(doc, "Ahmad Duais, Chief Executive: [BIO BULLETS: years in security, scale built, precinct/events/LP credentials, relevant governance. Half page.]")
bullet(doc, "[OPS MANAGER NAME], Operations Manager: [BIO: years, operational scope, incident command experience.]")
bullet(doc, "[NORTHLAND LEAD], Northland District Manager: [BIO or statement of the role profile being recruited, with target credentials: healthcare or emergency services background, Te Tai Tokerau connection.]")
bullet(doc, "[HEALTHCARE SPECIALIST], Healthcare Security Lead: [BIO: DHB/Health NZ or incumbent background, SPEC/SI credentials.]")
note(doc, "If the Northland and healthcare roles are not filled by submission, write the bio as a role profile with recruitment status and target date. A named person is worth real points; a credible funded role profile still scores. An empty promise does not.")

# ================= CAPACITY P6 =================
heading(doc, "CAPACITY - Q6: Transition plan, Employment Relations Act 2000 obligations, and capacity beyond transferred staff", 1)
para(doc, "Metro districts (Auckland, Counties Manukau, Waitemata): transition is deliberately light because these are supplementary on-demand services. Our mobilisation covers: standing up the dedicated Health NZ pool; completing all per-officer gates (COA verification, English evidence, Category A health screening started at preferred-supplier stage, Children's Act checks where applicable); Health NZ inductions and HSO Violence Reduction Training in the transition month after contract execution; the Auckland patrol and alarm-response run (vehicle, driver roster, SOPs) live from day one.", space_after=4)
para(doc, "Northland District: a full phased takeover plan, using the window to 1 July 2027 if required:", bold=True, space_after=2)
bullet(doc, "Phase 0 (award to contract execution): Northland District Manager and SOC lead appointments confirmed; local recruitment launched with [TE TAI TOKERAU PARTNER / planned partner]; engagement opened with incumbent staff (below).")
bullet(doc, "Phase 1 (months 1-2): hiring waves for Whangarei, Kaitaia, Kawakawa and Dargaville; COA processing for unlicensed local candidates; vetting and health screening running in parallel.")
bullet(doc, "Phase 2 (months 2-3): Health NZ inductions; HSO Violence Reduction Training as scheduled by the District; SPEC for any Mental Health assigned officers; SOC training on Avigilon, Gallagher and dispatch procedures; supervised handover shifts with the incumbent where Health NZ facilitates.")
bullet(doc, "Phase 3: full service commencement no later than 1 July 2027, with a [X]-week contingency buffer built into every phase.")
para(doc, "Employment Relations Act 2000 obligations:", bold=True, space_after=2)
bullet(doc, "We understand our obligations under the ERA, including Part 6A. We note that security officers are not a category listed in Schedule 1A, so no automatic statutory transfer right applies; our plan and pricing therefore do not depend on inheriting the incumbent workforce.")
bullet(doc, "Our intention nonetheless is continuity: subject to our screening standards, we will offer employment on ASG terms to suitable incumbent officers who wish to continue in their roles, through a fair and ERA-compliant process. In Northland this is also the most practical staffing route and the best outcome for local employment.")
bullet(doc, "Health NZ has confirmed that incumbent workforce information is not available during this process; our declaration and pricing reflect our own enquiries, and we will not seek additional payment for employment-law compliance costs, consistent with RFP section 4.2.")
para(doc, "Capacity beyond transferred staff: the recruitment pipeline in Q2 runs independently of any incumbent transfers, so the Pool for each district is filled whether or not incumbent staff accept offers. Transferees accelerate readiness; they are not load-bearing in the plan.", space_after=4)

# ================= H&S H1 =================
heading(doc, "HEALTH & SAFETY - Q1: Health and safety management practices, promoting and monitoring worker, patient, staff and public safety", 1)
para(doc, "Our health and safety system is built for a workforce whose daily job involves other people's worst moments. It covers our officers first, and through them the patients, Health NZ staff and public they protect.", space_after=4)
bullet(doc, "Framework: a current H&S policy under the Health and Safety at Work Act 2015 (copy attached to this response), site-level hazard and risk registers, documented safe work procedures for high-risk tasks (physical intervention, lone work, night work), and PPE issue and replacement standards. [CERTIFICATION STATUS: SiteWise/Impac/other prequalification, state what is real.]")
bullet(doc, "Officer due diligence: health and safety performance is a standing director agenda item; Ahmad Duais holds PCBU officer responsibility and reviews incident and hazard trends [FREQUENCY].")
bullet(doc, "Worker participation: [REAL MECHANISM: H&S representative, toolbox meetings at frequency, open hazard reporting channel]. Officers are trained that stopping unsafe work is an obligation, not an option.")
bullet(doc, "Protection of patients, staff and the public: our intervention doctrine is de-escalation first, least restrictive force, and clinical deference, in a hospital the clinical team leads and our officers support. Dynamic risk assessment is trained and expected before every physical intervention.")
bullet(doc, "Overlapping duties: we acknowledge the HSWA overlapping duties with Health NZ as PCBUs sharing a workplace, and commit to consulting, coordinating and cooperating, including joint site risk reviews with each District Security Manager and participation in Health NZ monitoring and audits.")
bullet(doc, "Monitoring: hazard reports, near misses, incidents and training currency are tracked in [SYSTEM]; trends are reviewed monthly and reported to clients as agreed.")

# ================= H&S H2 =================
heading(doc, "HEALTH & SAFETY - Q2: Fatigue management, rest periods, and secondary employment", 1)
para(doc, "Fatigue is a primary hazard in 24/7 security work and we manage it as one:", space_after=4)
bullet(doc, "Hard rostering rules enforced in Deputy: maximum shift length [X] hours; maximum [X] hours per week; minimum [X] hours between shifts; maximum [X] consecutive night shifts. The rostering system flags violations and they require [APPROVAL LEVEL] to override, with the reason recorded. [STATE REAL RULES.]")
bullet(doc, "Night and lone-worker welfare: scheduled welfare checks for lone posts [FREQUENCY/METHOD], which we would apply to single-officer hospital posts.")
bullet(doc, "Secondary employment: every officer declares secondary employment at engagement and re-declares [annually / on change]. Where secondary work is declared, combined hours are assessed against our fatigue limits before rosters are confirmed. Officers are contractually obliged to disclose changes, and supervisors are trained to spot fatigue indicators (lateness, presentation, alertness) and act on them.")
bullet(doc, "Within the wider ASG group, rosters are visible in a single system, so an officer cannot accumulate unsafe combined hours across group engagements undetected.")
bullet(doc, "Fitness for duty: a fatigued officer is stood down without blame and the shift is backfilled through the normal process. We treat a fatigue stand-down as the system working, not failing.")
bullet(doc, "We note the Waikato District schedule's restriction on secondary employment for fully deployed officers and will apply equivalent transparency for any Health NZ dedicated full-time roles in the districts we are bidding.")

# ================= H&S H3 =================
heading(doc, "HEALTH & SAFETY - Q3: Logging standard and notifiable incidents, and reporting out of the system", 1)
para(doc, "Every incident is logged digitally, on shift, by the officer involved:", space_after=4)
bullet(doc, "Platform: [PLATFORM NAME/DESCRIPTION]. Structured incident records with category, location, time, people involved, narrative, actions taken and outcome; photo attachment; supervisor review workflow. Records are timestamped and tamper-evident.")
bullet(doc, "Classification: standard incidents versus notifiable events under HSWA are distinguished at logging. A notifiable event triggers an immediate escalation workflow: scene preservation, duty manager notification, WorkSafe notification where required, and client notification. This matches the contract KPI: notifiable incidents reported to the District Security Manager within 24 hours, all others within 48; our standard practice is same-shift reporting.")
bullet(doc, "Health information discipline: for hospital sites, reports are written to Health Information Privacy Code standards, recording what is necessary for safety and legal purposes without gratuitous patient detail.")
bullet(doc, "Reporting out: the system produces daily, weekly and monthly summaries by category and site, client-accessible dashboards or scheduled exports, and automatic notification of completed incident reports to nominated client staff. We note the Northland schedule requires the supplier to provide exactly this capability (incident system with Health NZ data access, automatic notifications, and daily, weekly and monthly category summaries); it is our existing operating model, not a build.")
bullet(doc, "The same data drives prevention: incident trends feed training content and site SOP reviews (see Q5).")

# ================= H&S H4 =================
heading(doc, "HEALTH & SAFETY - Q4: Force majeure response, business continuity and emergency/risk management", 1)
bullet(doc, "Business continuity plan: documented, covering loss of key staff, rostering or communications system outage, vehicle loss, fuel or transport disruption, severe weather, and pandemic restrictions. Reviewed at least every twelve months and after any activation [LAST REVIEW DATE]; the contract requirement for BCP review evidence in reporting is noted and will be met.")
bullet(doc, "Communications redundancy: [REAL ARRANGEMENTS: secondary comms channel, printed contact trees, offline roster extracts].")
bullet(doc, "Emergency surge: the tiered activation model in Capacity Q3, including recall lists and cross-district support. In a declared emergency we accept the contract's best-endeavours obligation, including sourcing personnel from other regions, and our event heritage means rapid large-scale stand-up is demonstrated practice: [EXAMPLE with numbers, confirm disclosure].")
bullet(doc, "Pandemic and outbreak operation: infection-control PPE compliance under clinical direction is already part of our induction for healthcare settings; our Category A screening register means officer immunity status is known and exclusions during outbreak events can be managed quickly with backfill from the pool.")
bullet(doc, "Risk management: a maintained company risk register, scenario exercises [FREQUENCY, e.g. annual BCP tabletop], and post-incident reviews after every significant activation, with actions tracked to closure.")

# ================= H&S H5 =================
heading(doc, "HEALTH & SAFETY - Q5: Measuring and continuously improving health and safety performance", 1)
bullet(doc, "Leading indicators: training currency percentage, hazard and near-miss reports per month (we target more reporting, not less), site inspections completed, toolbox attendance. [CURRENT FIGURES if available.]")
bullet(doc, "Lagging indicators: recordable injuries, lost-time injuries, assault-related injuries per 1,000 deployed hours, and WorkSafe notifiable events [FIGURES for last 12 months; zero is a fine answer if true and verifiable].")
bullet(doc, "Review cycle: monthly operational H&S review; quarterly trend analysis shared with clients (and at Health NZ QBRs under this contract); annual full system review against the HSWA and current WorkSafe guidance.")
bullet(doc, "Improvement loop: every incident and near miss generates review actions with owners and due dates in [SYSTEM]; training content is refreshed from our own incident data so improvement is grounded in what actually happens on our sites.")
bullet(doc, "Example: [ONE REAL IMPROVEMENT with before/after, e.g. a change to intervention procedure or rostering after incident analysis, with measured effect. If none is citable, delete rather than invent.]")
bullet(doc, "Under this contract we would add healthcare-specific measures from commencement: patient-watch incident rates, de-escalation resolution percentage, and training currency against HSO Violence Reduction and SPEC revalidation dates, reported per district.")

doc.save(sys.argv[1] if len(sys.argv) > 1 else "HNZ_RFP_0018709_Draft_Answers_Part2.docx")
print("Written:", sys.argv[1] if len(sys.argv) > 1 else "HNZ_RFP_0018709_Draft_Answers_Part2.docx")
