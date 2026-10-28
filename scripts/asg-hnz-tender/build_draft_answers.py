#!/usr/bin/env python3
"""HNZ RFP 0018709 — draft answers: Economic Benefits Q1-Q3 + Capability Q1/Q7.

Outputs a DOCX with [PLACEHOLDERS] highlighted yellow for Ahmad's numbers.
Copy-paste ready for Response Form V3 blue sections. Dash-free body text.

Usage: .venv/bin/python build_draft_answers.py <out.docx>
"""
import re, sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_COLOR_INDEX

NAVY = RGBColor(0x1B, 0x2A, 0x4A)

def add_runs(p, text, size=10.5):
    for part in re.split(r"(\[[^\]]+\])", text):
        if not part:
            continue
        run = p.add_run(part)
        run.font.size = Pt(size)
        run.font.name = "Calibri"
        if part.startswith("["):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            run.bold = True

def para(doc, text, size=10.5, bold=False, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if bold:
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(size)
    else:
        add_runs(p, text, size)
    return p

def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    add_runs(p, text)

def heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = NAVY
    return h

def note(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("DRAFTING NOTE (delete before submission): " + text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)

doc = Document()

heading(doc, "HNZ RFP 0018709 - Draft Answers v1", 0)
para(doc, "Economic Benefits Q1-Q3 and Capability Q1/Q7, for Response Form V3 blue sections. Respondent: Absolute Security Group Limited. Districts: Auckland, Counties Manukau, Waitemata, Northland. Yellow [PLACEHOLDERS] need Ahmad's numbers or a disclosure decision. Red notes are internal and must be deleted.", 10)

# ============ EB Q1 ============
heading(doc, "ECONOMIC BENEFITS - Q1: Engagement by District with local iwi, Pacifica Hauora and other community networks", 1)
note(doc, "Name an organisation ONLY if first contact has actually happened by submission date. Anything not yet contacted stays as a committed plan with a date, never as engagement underway.")

para(doc, "ASG is an Auckland-founded, New Zealand-owned security company whose workforce lives in the communities this contract serves. Our officers are predominantly recruited from South Auckland, West Auckland and Central Auckland, and our community engagement reflects where our people come from, not a corporate programme run at arm's length.", space_after=4)
para(doc, "Current engagement:", bold=True, space_after=2)
bullet(doc, "Community safety partnership working: through our city-centre precinct security programme we work alongside [PRECINCT CLIENT NAME, confirm disclosure], NZ Police, local businesses and social services on street-level safety, including referral and de-escalation of people in mental distress and rough sleepers, rather than enforcement-first responses.")
bullet(doc, "[INSERT any real current community engagement: sports sponsorship, marae connection, church/community group relationships of ASG or its staff, school or training programme involvement. If none, delete this bullet.]")
bullet(doc, "Our recruitment already runs through community channels: word of mouth within Pasifika and Maori whanau networks, [COMMUNITY JOB CHANNELS USED, e.g. local WINZ/MSD relationships, community job fairs], which is why our workforce mirrors the districts we would serve.")
para(doc, "Committed engagement plan for this contract (owner: Ahmad Duais, CEO; budget [$X] in year one):", bold=True, space_after=2)
bullet(doc, "Auckland and Waitemata Districts: we have [initiated contact with / scheduled a first meeting with] [IWI ORGANISATION, e.g. Ngati Whatua Orakei Whai Maia, ONLY IF CONTACTED, date] to establish a workforce pipeline partnership for healthcare security roles, including recruitment, pastoral support and cultural competency input into our training. [If no contact made: replace with: We commit to establishing a formal workforce pipeline partnership with a mana whenua organisation in the Auckland region, with first engagement scheduled by (DATE) and a signed partnership objective by contract commencement.]")
bullet(doc, "Counties Manukau District: engagement with [SOUTH AUCKLAND PACIFIC HAUORA / COMMUNITY PROVIDER, ONLY IF CONTACTED, date] focused on recruiting and supporting Pasifika officers for the Middlemore patient-watch workforce, where cultural and language alignment with the patient population directly improves de-escalation outcomes.")
bullet(doc, "Northland District: our Northland mobilisation plan is built on local hiring, not importing Auckland staff. We [have contacted / will engage by (DATE)] [TE TAI TOKERAU IWI / TRAINING PROVIDER, ONLY IF CONTACTED] to recruit and train officers living in Whangarei, Kaitaia, Kawakawa and Dargaville for the scheduled posts at each hospital, timed to full readiness between 1 April and 1 July 2027. This creates an estimated [N] permanent local roles in towns with few skilled-work opportunities.")
bullet(doc, "Governance: progress against this plan is reported in our Quarterly Business Reviews with Health NZ, with named counterparts, milestones and outcomes, so it remains accountable rather than aspirational.")

# ============ EB Q2 ============
heading(doc, "ECONOMIC BENEFITS - Q2: Building a diverse and inclusive security workforce reflective of local communities", 1)
para(doc, "Our workforce today:", bold=True, space_after=2)
bullet(doc, "[X]% of our [N] active officers identify as Maori, [X]% as Pasifika, [X]% as other ethnic communities; our team speaks [N] languages including [te reo Maori / Samoan / Tongan / Hindi / etc].")
bullet(doc, "[X]% of officers live in South or West Auckland; our Counties Manukau and Waitemata deployments are staffed overwhelmingly by people who live in those districts.")
bullet(doc, "[X]% of our workforce is female [adjust/delete to reflect reality; add any age-diversity point if useful].")
note(doc, "These percentages must be real and reproducible from PayHero/Deputy records. Evaluators can test them at reference checks.")
para(doc, "What makes it work (current practice, not aspiration):", bold=True, space_after=2)
bullet(doc, "Recruitment through community and whanau networks rather than generic job boards only, which produces candidates with standing in their own communities, a genuine asset in a hospital where patients and whanau respond better to officers who look and sound like their community.")
bullet(doc, "Structured English language support and assessment against the Immigration NZ Skilled Migrant Category standard, so officers from migrant communities are supported to meet Health NZ's requirement rather than screened out by it.")
bullet(doc, "Inclusive rostering: [DESCRIBE real practices: religious observance accommodation, fair allocation of penal-rate shifts, parental flexibility. Delete what is not true.]")
para(doc, "Plan for the Health NZ pool:", bold=True, space_after=2)
bullet(doc, "We will build the dedicated Health NZ officer pool to reflect each district served: Counties Manukau mirroring South Auckland's Pasifika and Maori communities, Waitemata mirroring West Auckland, and Northland recruited locally in Te Tai Tokerau with a target of [X]% Maori representation, supported by the iwi partnership described in Q1.")
bullet(doc, "Cultural competency is built into induction and ongoing training: [DESCRIBE current content, e.g. tikanga-informed engagement, use of te reo greetings, working with whanau groups; or state the committed plan with provider and date].")
bullet(doc, "Measurement: workforce composition versus district demographics reported annually at QBRs.")

# ============ EB Q3 ============
heading(doc, "ECONOMIC BENEFITS - Q3: Strengthening New Zealand's capability and capacity in healthcare security services", 1)
para(doc, "ASG's contribution to national healthcare security capability rests on three commitments.", space_after=4)
para(doc, "1. A skilled career pathway, not a doorway job.", bold=True, space_after=2)
bullet(doc, "Every ASG officer enters a structured development pathway: frontline officer, then specialist capability (covert and loss prevention tradecraft, control room and CCTV operation, event command), then senior officer and supervision. [N] of our current senior staff began as frontline guards, including [EXAMPLE ROLE/NAME if willing].")
bullet(doc, "We invest [X hours / $X] per officer per year in training beyond licensing minimums, anchored by our in-house field training manual and scenario-based de-escalation training, and we commit to aligning this programme with the Health NZ Healthcare Security Officer Violence Reduction Training framework from contract commencement.")
para(doc, "2. Bringing covert and loss prevention tradecraft into healthcare security.", bold=True, space_after=2)
bullet(doc, "ASG's group operates one of New Zealand's most active retail loss prevention programmes, with more than [180+] trained officers active nationally and over [$1M] of verified recoveries since July 2025 for national retail clients. The underlying skills, behavioural observation, early identification of escalating individuals, covert patrol, evidence-quality reporting and lawful intervention, barely exist in New Zealand hospital security today.")
bullet(doc, "Under this contract those skills would be delivered by ASG-employed officers and offered to Districts as an additional capability, subject to District Security Manager approval and Health NZ privacy settings: covert presence in car parks and public zones, identification of repeat trespassers and individuals targeting staff, and protection of hospital equipment and property. This builds a new national capability in the healthcare security workforce, and the officers we train in it remain in the New Zealand industry.")
para(doc, "3. Local capacity, New Zealand ownership.", bold=True, space_after=2)
bullet(doc, "ASG is 100% New Zealand owned; every dollar of margin stays in New Zealand.")
bullet(doc, "This contract creates an estimated [N] new roles across the four districts, including [N] in Northland where we commit to local hiring and training in Whangarei, Kaitaia, Kawakawa and Dargaville (see Q1).")
bullet(doc, "Healthcare-specific uplift: we are [recruiting / have engaged] a healthcare security specialist [STATUS: name/role if secured by submission] to lead our hospital capability, develop our patient watch operating procedures, and embed SPEC and Safety Intervention practice across the Health NZ pool, so the sector gains an additional credible supplier rather than deeper dependence on the incumbents.")

# ============ Capability Q1 ============
heading(doc, "CAPABILITY - Q1: Experience in a healthcare (or similarly intensive) environment", 1)
note(doc, "Strategy: honest about no hospital contracts, then translation with HARD numbers, then the dated readiness plan. District SMEs will anchor on incumbent-grade answers; the numbers and the referee are what move this from 3-4 to 5-7.")
para(doc, "ASG does not currently hold hospital security contracts. We ask the evaluation panel to assess us on what the RFP invites: a detailed demonstration that our experience in similarly intensive environments translates to healthcare, together with a dated plan for healthcare-specific readiness at contract commencement.", space_after=4)
para(doc, "Where our experience comes from:", bold=True, space_after=2)
bullet(doc, "City-centre entertainment precinct security, [PRECINCT CLIENT NAME, confirm disclosure], [N] nights per week since [DATE]: our officers manage intoxication, mental distress, drug-affected behaviour, verbal aggression and physical violence in crowded public settings every week. In the last 12 months: [N] recorded incidents, [X]% resolved through verbal de-escalation without physical intervention, [N] involving persons in mental distress, average response time [X] minutes. These figures come from our incident reporting platform and our referee for this engagement will verify them.")
bullet(doc, "National retail loss prevention (group programme, ASG-employed delivery for this contract): over [180+] active officers, [$1M+] verified recoveries since July 2025, daily lawful interventions with offenders who are frequently aggressive, drug-affected or in crisis. The core healthcare security skill, controlling a volatile person in a public space full of bystanders while staying inside the law and preserving dignity, is our daily work.")
bullet(doc, "Major event and surge operations, including [CRL opening events / other, confirm disclosure]: rapid stand-up of large deployments, command structure, coordination with Police and emergency services.")
bullet(doc, "24/7 operations: a staffed operations capability dispatching and supervising officers around the clock across Auckland, with live response data: [average urgent response time X minutes across Auckland metro, if available].")
para(doc, "How it translates (requirement to evidence):", bold=True, space_after=2)
bullet(doc, "Managing mental distress, cognitive impairment and substance-related harm: [N] mental-distress incidents handled in 12 months in precinct work; de-escalation-first doctrine with physical intervention as last resort ([X]% verbal resolution rate).")
bullet(doc, "Working alongside frontline staff with patients, whanau and visitors present: our officers work every shift alongside venue staff, retail staff and the public; customer service presentation is contractual at our current clients, as it is in this RFP.")
bullet(doc, "Low-risk to high-risk settings: the same officer moves between a quiet foyer and a violent ejection within minutes in precinct work, matching the hospital reality of moving between reception presence and an ED crisis.")
bullet(doc, "Reporting discipline: evidence-quality incident reports on a digital platform with client access, matching Health NZ's 24/48-hour incident reporting KPI.")
para(doc, "Healthcare readiness plan (complete before service commencement, 1 April 2027; Northland by 1 July 2027):", bold=True, space_after=2)
bullet(doc, "[MONTH]: healthcare security specialist [appointed / appointment completed]; patient watch and clinical-environment operating procedures drafted with their lead.")
bullet(doc, "Transition month: all assigned officers complete Health NZ inductions and the HSO Violence Reduction Training at the level set by each District Security Manager; SPEC (4-day) for any officer assigned to Mental Health facilities.")
bullet(doc, "Pre-assignment: every officer holds a full Property Guard COA with NZQA US 27364, US 27360 and US 27361, verified English language evidence, and completed infectious disease screening per Attachments W and X.")
para(doc, "District note: this experience base serves all four districts bid. For Northland, delivery additionally depends on the local workforce and Security Operations Centre plan described in Capacity, led by [NORTHLAND LEAD ROLE/NAME].", space_after=4)

# ============ Capability Q7 ============
heading(doc, "CAPABILITY - Q7 (Question 1.g): Approach to other services outside Scheduled and On-Demand security officer services", 1)
para(doc, "Our approach by district, covering the services in sections 3 to 8 of each District Schedule:", space_after=4)
para(doc, "Auckland District:", bold=True, space_after=2)
bullet(doc, "24/7 Alarm Response across the designated facilities (Auckland City Hospital, Greenlane Clinical Centre, Buchanan Rehabilitation Centre) and the 12 ARDS community dental clinics: delivered by a dedicated ASG response vehicle and driver-qualified officers operating from our Auckland base, with response procedures agreed in SOPs. We note Attachment F lists no Health NZ vehicle; our pricing approach will include supplier vehicle provision.")
bullet(doc, "Mobile and foot patrols at the six listed locations, seven days including bed-down checks at the mental health and rehabilitation facilities and night checks at Mankaai CMHC: delivered on the same vehicle run, by officers trained for mental health settings.")
bullet(doc, "Emergency support: surge deployment from our Auckland pool under our 24/7 operations escalation procedures.")
para(doc, "Counties Manukau District:", bold=True, space_after=2)
bullet(doc, "Alarm response when dispatched by Health NZ personnel; all assigned officers trained to working knowledge of CCTV operation.")
bullet(doc, "Transportation of cash where required: [CONFIRM broker: cash-in-transit insurance extension] with two-officer procedures.")
bullet(doc, "Foot patrols as required across the nine district sites.")
para(doc, "Waitemata District:", bold=True, space_after=2)
bullet(doc, "Alarm monitoring and response in line with location manager requirements; CCTV monitoring competency and operation of the onsite Gallagher access control system are part of our assigned officers' induction and verification before deployment.")
bullet(doc, "Mobile patrols using the Health NZ supplied vehicle; we verify every assigned officer holds at least a Restricted licence.")
bullet(doc, "Specialist Mental Health Services deployments at short notice, from a bench of SPEC-trained officers maintained for this purpose.")
para(doc, "Northland District:", bold=True, space_after=2)
bullet(doc, "Security Operations Centre: we will staff the Whangarei SOC with a dedicated Operations Controller capability, [NAMED SOC LEAD / recruitment status], trained across alarm monitoring, the district's 300+ camera CCTV network (Avigilon), Gallagher access control, dispatch of our officers, log management and urgent information requests. Controller succession and relief are built into the Northland pool so SOC capability never rests on one person.")
bullet(doc, "Off-site night patrols at 3 Grant Street Te Kamo and Manaia House (Rathbone Street), two checks nightly each, plus mobile attendance at building alarms at those locations: delivered by non-hospital officers with an ASG patrol vehicle based in Whangarei.")
bullet(doc, "On-site mobile and foot patrols at all four hospital campuses per SOP.")
para(doc, "Additional capability offered across all districts (subject to District Security Manager approval): covert and plain-clothes officers for car parks and public non-clinical areas, targeting vehicle crime, theft of hospital equipment and property, and identification of repeat trespassers and individuals targeting staff. This draws on our established loss prevention methodology and is delivered by ASG-employed officers within Health NZ privacy requirements. All services above feed a single incident reporting platform with weekly service reporting available to each District Security Manager.", space_after=4)

doc.save(sys.argv[1] if len(sys.argv) > 1 else "HNZ_RFP_0018709_Draft_Answers.docx")
print("Written:", sys.argv[1] if len(sys.argv) > 1 else "HNZ_RFP_0018709_Draft_Answers.docx")
