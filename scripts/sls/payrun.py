#!/usr/bin/env python3
"""SLS RTC pay-run: unified timesheet xlsx -> PayHero import CSV.
Implements ~/.claude/commands/adjust-paycycle-timesheet.md rules."""
import sys, os, csv, re, shutil
from datetime import datetime
import openpyxl

EXCLUDE = {"faiez idais", "zac shannon"}

def norm_name(first, last):
    name = f"{first or ''} {last or ''}"
    name = re.sub(r"\s+", " ", name).strip()
    return " ".join(w.capitalize() for w in name.split())

def parse_hm(v):
    if v is None: return None
    s = str(v).strip()
    if not s: return None
    m = re.match(r"^(\d{1,2}):(\d{2})", s)
    if not m: return None
    return int(m.group(1)) * 60 + int(m.group(2))

def fmt_time(v):
    """Format HH:MM -> H:MM (strip leading zero on hour)."""
    m = re.match(r"^(\d{1,2}):(\d{2})", str(v).strip())
    return f"{int(m.group(1))}:{m.group(2)}"

def process(xlsx, out_dir):
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb["All Employees"] if "All Employees" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]
    idx = {h: i for i, h in enumerate(hdr)}
    c_first, c_last, c_type = 0, 1, 2
    c_sdate, c_in, c_out = idx["Start Date"], idx["In"], idx["Out"]

    shifts = {}          # (emp, dateobj) -> list of (start_min, end_min, in_str, out_str)
    dropped_excl, dropped_leave, dropped_ph, dropped_zero, skipped = {}, [], [], [], []
    flags = []

    for ri, r in enumerate(rows[1:], start=2):
        first, last, typ = r[c_first], r[c_last], r[c_type]
        emp = norm_name(first, last)
        if not emp:
            continue
        if emp.lower() in EXCLUDE:
            dropped_excl[emp] = dropped_excl.get(emp, 0) + 1
            continue
        tl = (str(typ) if typ else "").lower()
        if "public holiday" in tl:
            dropped_ph.append((emp, typ)); continue
        if "leave" in tl or "holiday" in tl:
            dropped_leave.append((emp, typ)); continue
        sdate, vin, vout = r[c_sdate], r[c_in], r[c_out]
        if not isinstance(sdate, datetime):
            if vin or vout: skipped.append((ri, emp, "no date"))
            continue
        smin, emin = parse_hm(vin), parse_hm(vout)
        if smin is None or emin is None:
            skipped.append((ri, emp, "missing in/out")); continue
        if smin == emin:
            dropped_zero.append((emp, sdate.date())); continue
        dur = (emin - smin) / 60.0
        if dur < 0: dur += 24  # overnight
        if dur > 14:
            flags.append((emp, sdate.date(), fmt_time(vin), fmt_time(vout), round(dur,2)))
        key = (emp, sdate.date())
        shifts.setdefault(key, []).append((smin, emin, fmt_time(vin), fmt_time(vout)))

    # overlap detection
    overlaps = []
    for (emp, d), lst in shifts.items():
        s = sorted(lst)
        for i in range(len(s)):
            for j in range(i+1, len(s)):
                if s[i][0] < s[j][1] and s[j][0] < s[i][1]:
                    overlaps.append((emp, d, s[i][2:], s[j][2:]))

    # Pass 2: breaks + build output rows
    out = []
    for (emp, d), lst in shifts.items():
        s = sorted(lst)
        total_min = sum(e-st if e>=st else (e-st+1440) for st,e,_,_ in s)
        break_dur = {}  # index into s -> break value
        if total_min >= 360:
            gap = 0
            for i in range(1, len(s)):
                g = s[i][0] - s[i-1][1]
                if g > 0: gap += g
            if gap < 30:
                shortfall = 30 - gap
                longest = max(range(len(s)), key=lambda i: (s[i][1]-s[i][0]))
                break_dur[longest] = round(shortfall/60.0, 2)
        for i,(st,e,ins,outs) in enumerate(s):
            out.append((emp, "Default Work", d.strftime("%d/%m/%y"), ins, outs,
                        "", "", "", (break_dur.get(i, "") if break_dur.get(i) else ""), ""))

    # sort: employee asc, date desc within employee
    def keyf(row):
        d = datetime.strptime(row[2], "%d/%m/%y")
        # also stable secondary by start time desc within day to mirror prior outputs
        return (row[0].lower(), -d.toordinal())
    out.sort(key=lambda row: (row[0].lower(), -datetime.strptime(row[2], "%d/%m/%y").toordinal(),
                              -parse_hm(row[3])))

    os.makedirs(out_dir, exist_ok=True)
    dst_xlsx = os.path.join(out_dir, os.path.basename(xlsx))
    if os.path.abspath(xlsx) != os.path.abspath(dst_xlsx):
        shutil.copy2(xlsx, dst_xlsx)
    csv_path = os.path.join(out_dir, "PayHero_-_Timesheet_Import_Template.csv")
    with open(csv_path, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Employee","Work Name","Date","Start Time","End Time",
                    "Employee Email","Notes","Duration","Break Duration","Units"])
        w.writerows(out)

    emps = sorted({r[0] for r in out})
    dates = sorted({r[2] for r in out})
    breaks = sum(1 for r in out if r[8])
    print(f"\n=== {os.path.basename(out_dir)}  ({os.path.basename(xlsx)}) ===")
    print(f"  Output rows: {len(out)} | employees: {len(emps)} | date range: {dates[0]}..{dates[-1]}")
    print(f"  Break deductions applied: {breaks}")
    if dropped_excl: print(f"  Excluded: " + ", ".join(f"{k}({v})" for k,v in dropped_excl.items()))
    print(f"  Dropped — leave: {len(dropped_leave)}, public holiday: {len(dropped_ph)}, zero-duration: {len(dropped_zero)}")
    if dropped_ph:
        php = {}
        for e,t in dropped_ph: php[e]=php.get(e,0)+1
        print(f"    PH by emp: " + ", ".join(f"{k}({v})" for k,v in php.items()))
    if dropped_leave:
        lvp = {}
        for e,t in dropped_leave: lvp[f"{e}:{t}"]=lvp.get(f"{e}:{t}",0)+1
        print(f"    Leave: " + ", ".join(f"{k}({v})" for k,v in lvp.items()))
    if skipped: print(f"  Skipped (data issues): {len(skipped)} -> {skipped[:8]}")
    if flags: print(f"  ⚠ Shifts >14h: {flags}")
    if overlaps:
        print(f"  ⚠ OVERLAPS ({len(overlaps)}):")
        for o in overlaps: print(f"      {o[0]} {o[1]} {o[2]} vs {o[3]}")
    return csv_path

if __name__ == "__main__":
    process(sys.argv[1], sys.argv[2])
