#!/usr/bin/env python3
"""WW data-integrity validator — locks in the 5 fixes from the 2 Jul audit
(remediated 13 Jul 2026). Run any time; exits non-zero on any failure.

Checks:
  1. Canonical FY26 master: YTD Total Loss (G141) equals the carry-forward
     (G129+B141+C141) and never reads below the prior block's YTD.
  2. Canonical FY26 master: zero #REF! cells.
  3. Stale root masters stay neutralised (no un-prefixed root master exists).
  4. truth.db ww_deterrents: no gap wider than GAP_DAYS within FY26.
  5. truth.db: zero 'Hastings Woolworths' rows in ww_incidents/ww_deterrents.
  6. (bonus) FY26 consolidated GIR: zero junk-signature rows, zero dup IDs.
"""
import os, sys, sqlite3, datetime, collections
import openpyxl

CANON = os.path.expanduser("~/Desktop/Woolworths /25th May 2025/Woolworths Master Data.xlsx")
ROOT_DIR = os.path.expanduser("~/Desktop/Woolworths ")
DB = os.path.expanduser("~/Desktop/Dev/asg-ops-hub/data/truth.db")
CONSOLIDATED = os.path.expanduser(
    "~/Desktop/SLS-RTC/Operations/Analytics/data/latest/gir-woolworths-FY26-consolidated.xlsx")
GAP_DAYS = 6            # widest tolerated silence in deterrent events
FY_S, FY_E = datetime.date(2025, 7, 1), datetime.date(2026, 6, 30)

failures = []
def check(ok, label, detail=""):
    print(f"  {'PASS' if ok else 'FAIL'}  {label}" + (f" — {detail}" if detail else ""))
    if not ok:
        failures.append(label)

print("== 1+2. Canonical FY26 master ==")
wb = openpyxl.load_workbook(CANON, read_only=True, data_only=True)
ws = wb["Total Master"]
g129, b141, c141, g141 = (ws["G129"].value, ws["B141"].value or 0,
                          ws["C141"].value or 0, ws["G141"].value)
expected = round(float(g129) + float(b141) + float(c141), 2)
check(g141 is not None and abs(float(g141) - expected) < 0.01,
      "YTD carry-forward intact", f"G141={g141} expected={expected}")
check(g141 is not None and float(g141) >= float(g129),
      "YTD never drops below prior block", f"G141={g141} G129={g129}")
refs = sum(1 for row in ws.iter_rows(min_row=1, max_row=250) for c in row
           if isinstance(c.value, str) and "#REF" in c.value)
check(refs == 0, "zero #REF! cells", f"found {refs}")
wb.close()

print("== 3. Root masters neutralised ==")
bad = [f for f in os.listdir(ROOT_DIR)
       if f in ("Woolworths Master Data.xlsx", "Woolworths Master Data BACKUP.xlsx")]
check(not bad, "no un-prefixed FY26 root master", str(bad))

print("== 4+5. truth.db ==")
db = sqlite3.connect(DB)
dates = [r[0] for r in db.execute(
    "SELECT DISTINCT date(event_date) FROM ww_deterrents "
    "WHERE date(event_date) BETWEEN ? AND ? ORDER BY 1",
    (FY_S.isoformat(), FY_E.isoformat()))]
worst, prev = 0, None
for ds in dates:
    d = datetime.date.fromisoformat(ds)
    if prev: worst = max(worst, (d - prev).days)
    prev = d
check(worst <= GAP_DAYS, f"deterrent gaps <= {GAP_DAYS} days", f"worst gap {worst}d")
h_inc = db.execute("SELECT COUNT(*) FROM ww_incidents WHERE store LIKE '%Hastings%'").fetchone()[0]
h_det = db.execute("SELECT COUNT(*) FROM ww_deterrents WHERE store LIKE '%Hastings%'").fetchone()[0]
check(h_inc == 0 and h_det == 0, "zero Hastings mis-tags", f"inc={h_inc} det={h_det}")
db.close()

print("== 6. FY26 consolidated GIR ==")
wb = openpyxl.load_workbook(CONSOLIDATED, read_only=True, data_only=True)
ws = wb.worksheets[0]
rows = ws.iter_rows(values_only=True)
hdr = [str(h).strip() if h is not None else "" for h in next(rows)]
prod_i = next(i for i, h in enumerate(hdr) if h.startswith("General Product Categories"))
ids = collections.Counter(); junk = 0
for r in rows:
    ids[str(r[0])] += 1
    if (str(r[prod_i]).strip() if r[prod_i] else "") == "Image":
        junk += 1
wb.close()
dupes = sum(v - 1 for v in ids.values() if v > 1)
check(junk == 0, "zero junk-signature rows", f"found {junk}")
check(dupes == 0, "zero duplicate submission IDs", f"found {dupes}")

print()
if failures:
    print(f"FAILED: {len(failures)} check(s): {failures}")
    sys.exit(1)
print("ALL CHECKS PASSED")
