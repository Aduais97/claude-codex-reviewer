#!/usr/bin/env python3
"""Rebuild gir-woolworths-FY26-consolidated.xlsx from the authoritative full export.

WHY (diagnosed 13 Jul 2026):
The old consolidated file was assembled by POSITIONAL concatenation of era exports
(91-col fy.xlsx + 79/75/73-col weekly exports after the Jotform form changed
~18 May 2026). Result: ~69 post-18-May rows misaligned (junk signature: "Image"
in products, manager titles in recovery, $0 values), 154 era-overlap duplicate
rows, and 76 older misaligned orphans (Jul-Aug 2025, since deleted from Jotform).

FIX: rebuild from ONE source — the full-year Jotform export downloaded 2 Jul 2026
(91 cols, every submission re-rendered under the complete current field set,
zero junk signatures, includes 43 late-filed rows the weeklies missed).

Guarantees enforced (hard failures, not warnings):
  1. Header fingerprint: all REQUIRED_FIELDS present by name.
  2. FY26 date filter (1 Jul 2025 - 30 Jun 2026) on incident date.
  3. Zero duplicate submission IDs.
  4. Zero junk-signature rows (products == "Image", or recovery value outside
     the known set).
  5. Old file backed up, dropped rows preserved to an orphans file - nothing
     silently discarded.

Usage:
  rebuild_ww_fy26_consolidated.py [--source <export.xlsx>]
Defaults to the verified 2-Jul full export in ~/Downloads.
"""
import os, re, sys, shutil, datetime, collections
import openpyxl

LATEST = os.path.expanduser("~/Desktop/SLS-RTC/Operations/Analytics/data/latest")
TARGET = os.path.join(LATEST, "gir-woolworths-FY26-consolidated.xlsx")
DEFAULT_SOURCE = os.path.expanduser(
    "~/Downloads/Woolworths _ Guard Incident Report_c825cc65-4f04-47bf-b60c-57e23f7026ea.xlsx")
FY_S, FY_E = datetime.date(2025, 7, 1), datetime.date(2026, 6, 30)

REQUIRED_FIELDS = [
    "#", "Full name", "Date and Time of Incident - Date", "Region of site",
    "Select Woolworths Site", "Incident Severity", "Incident Category",
    "Type of Deterrent", "Were the police informed?",
    "General Product Categories Targeted", "Type of recovery",
    "Recovered Reciept Value", "Estimated Stolen Value",
]
VALID_RECOVERY = {"", "Full Recovery", "Partial Recovery", "No Recovery"}


def parse_date(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if v:
        try:
            return datetime.date.fromisoformat(str(v)[:10])
        except ValueError:
            return None
    return None


def main():
    source = DEFAULT_SOURCE
    if "--source" in sys.argv:
        source = sys.argv[sys.argv.index("--source") + 1]
    stamp = datetime.date.today().isoformat()

    # ---------- read source ----------
    wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    hdr = [str(h).strip() if h is not None else "" for h in next(rows)]

    # Guarantee 1 - header fingerprint
    missing = [f for f in REQUIRED_FIELDS if not any(h == f or h.startswith(f) for h in hdr)]
    if missing:
        sys.exit(f"FATAL: source header missing required fields: {missing}\n"
                 f"Form schema has drifted again - do NOT merge positionally. "
                 f"Update REQUIRED_FIELDS mapping deliberately.")

    def col(name):
        for i, h in enumerate(hdr):
            if h == name or h.startswith(name):
                return i
        raise KeyError(name)

    di, prod_i, rec_i = (col("Date and Time of Incident - Date"),
                         col("General Product Categories Targeted"),
                         col("Type of recovery"))

    kept, ids, junk = [], collections.Counter(), []
    for r in rows:
        d = parse_date(r[di])
        if d is None or not (FY_S <= d <= FY_E):
            continue                       # Guarantee 2 - FY26 window
        ids[str(r[0])] += 1
        prod = str(r[prod_i]).strip() if r[prod_i] else ""
        rec = str(r[rec_i]).strip() if r[rec_i] else ""
        if prod == "Image" or rec not in VALID_RECOVERY:
            junk.append((str(r[0]), str(d), prod[:20], rec[:20]))
        kept.append(r)

    dupes = {k: v for k, v in ids.items() if v > 1}
    if dupes:                              # Guarantee 3
        sys.exit(f"FATAL: duplicate submission IDs in source: {list(dupes)[:10]}")
    if junk:                               # Guarantee 4
        sys.exit(f"FATAL: {len(junk)} junk-signature rows in source "
                 f"(misalignment persists): {junk[:5]}")

    # ---------- reconcile against old file ----------
    old_ids = set()
    old_rows = 0
    if os.path.exists(TARGET):
        owb = openpyxl.load_workbook(TARGET, read_only=True, data_only=True)
        ows = owb.worksheets[0]
        orows = ows.iter_rows(values_only=True)
        ohdr = [str(h).strip() if h is not None else "" for h in next(orows)]
        odi = next(i for i, h in enumerate(ohdr)
                   if h.startswith("Date and Time of Incident - Date"))
        old_data = []
        for r in orows:
            d = parse_date(r[odi])
            if d and FY_S <= d <= FY_E:
                old_rows += 1
                old_ids.add(str(r[0]))
                old_data.append(r)
        owb.close()

        # Guarantee 5 - backup + orphans
        backup = TARGET.replace(".xlsx", f".CORRUPT-BACKUP-{stamp}.xlsx")
        shutil.copy2(TARGET, backup)
        new_ids = set(ids)
        orphan_ids = old_ids - new_ids
        if orphan_ids:
            ob = openpyxl.Workbook()
            obs = ob.active; obs.title = "Orphans"
            obs.append(ohdr + [f"NOTE: not present in {os.path.basename(source)}"])
            for r in old_data:
                if str(r[0]) in orphan_ids:
                    obs.append(list(r))
            orphan_path = TARGET.replace(".xlsx", f".orphans-{stamp}.xlsx")
            ob.save(orphan_path)
        print(f"old file:    {old_rows} FY26 rows, {len(old_ids)} unique IDs "
              f"({old_rows - len(old_ids)} duplicates)")
        print(f"backup:      {os.path.basename(backup)}")
        if orphan_ids:
            print(f"orphans:     {len(orphan_ids)} rows preserved -> "
                  f"{os.path.basename(orphan_path)}")
        print(f"recovered:   {len(new_ids - old_ids)} IDs in source but "
              f"missing from old file (late-filed submissions)")

    # ---------- write rebuilt file ----------
    nb = openpyxl.Workbook(write_only=True)
    nws = nb.create_sheet("Entries")
    nws.append(hdr)
    kept.sort(key=lambda r: parse_date(r[di]) or FY_S, reverse=True)
    for r in kept:
        nws.append(list(r))
    nb.save(TARGET)
    wb.close()

    print(f"\nrebuilt:     {os.path.basename(TARGET)}")
    print(f"             {len(kept)} FY26 rows, {len(ids)} unique IDs, "
          f"0 duplicates, 0 junk-signature rows")
    print(f"source:      {os.path.basename(source)}")


if __name__ == "__main__":
    main()
