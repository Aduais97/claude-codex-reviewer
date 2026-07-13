#!/usr/bin/env python3
"""
The Warehouse (Takanini) — weekly covert LP report generator (SLS RTC).
Single-store variant of the Mitre 10 weekly report. Reads the weekly Guard
Incident Report + Recovered Packaging xlsx from the week folder and produces:
  1. <prefix>_LP_Report.pdf     — incidents + recovered packaging
  2. <prefix>_Email_Summary.docx — client email summary
(No store assessment in week 2 — that was the launch Store Hardening report.)

Recovery $ values are entered manually (the Warehouse forms capture no $ field),
keyed by the report "#" id.
"""
import sys, os, glob, base64, subprocess, datetime
import openpyxl

WEEK = sys.argv[1] if len(sys.argv) > 1 else os.path.expanduser("~/Desktop/WarehouseGroup/Week ending 28th June")
PERIOD = sys.argv[2] if len(sys.argv) > 2 else "22 – 28 June 2026"
PREFIX = "Warehouse_Takanini_Week_Ending_28_June"
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
LOGO = os.path.expanduser("~/Desktop/SLS-RTC/Operations/Analytics/assets/slsrtc_logo_hd.png")

# Manually-entered recovery values (forms capture no $ field), keyed by report "#".
RECOVERIES_BY_ID = {15: 98.95, 13: 48.00, 12: 121.92, 11: 52.97}   # product recovered at incidents
PKG_VALUES_BY_ID = {7: 186.91, 6: 121.92}                          # recovered-packaging receipt values

ACC="#1f3a5f"; SUB="#5b6472"; LINE="#dfe3e8"; RED="#b23a3a"; GREEN="#1d7a4d"; CARD="#f6f8fa"; AMBER="#c98a1a"

def find(key):
    g = glob.glob(os.path.join(WEEK, f"*{key}*.xlsx"))
    return g[0] if g else None

def rows(fp):
    ws = openpyxl.load_workbook(fp, data_only=True).active
    rr = list(ws.iter_rows(values_only=True))
    hdr = [("" if h is None else str(h).strip()) for h in rr[0]]
    return [{h: v for h, v in zip(hdr, r) if h} for r in rr[1:]
            if any(v is not None and str(v).strip() for v in r)]

def col(d, *cands):
    for k in d:
        for c in cands:
            if c.lower() in k.lower():
                return d[k]
    return None

def store_short(s):
    s = str(s or "")
    return "Takanini" if "takanini" in s.lower() else s.replace("The Warehouse", "").strip() or "Takanini"

def hhmm(t):
    if isinstance(t, (datetime.datetime, datetime.time)): return t.strftime("%H:%M")
    return str(t or "")[:5]

def ddmm(d):
    if isinstance(d, datetime.datetime): return d.strftime("%d/%m")
    return str(d or "")[:5]

def iid_of(d):
    try: return int(float(d.get("#")))
    except (TypeError, ValueError): return None

# ---------- parse incidents ----------
inc = []
for d in rows(find("Guard Incident")):
    cat = str(col(d, "Incident Category") or "").strip()
    cause = col(d, "Cause/Reason for Intervention")
    if cat == "Deterrent": cause = col(d, "Type of Deterrent") or "Deterred"
    elif cat == "Other":   cause = col(d, "If other") or "Other"
    iid = iid_of(d)
    inc.append(dict(
        iid=iid,
        store=store_short(col(d, "Select The Warehouse Site", "Select The Warehouse", "Site")),
        date=ddmm(col(d, "of Incident - Date")), time=hhmm(col(d, "of Incident - Time")),
        cat=cat, cause=str(cause or "—").strip(),
        vt=str(col(d, "Verbal Trespass") or "No").strip(),
        rec=RECOVERIES_BY_ID.get(iid, 0.0),
    ))
inc.sort(key=lambda x: (x["date"], x["time"]))

# ---------- parse packaging ----------
pkg = []
for d in rows(find("Recovered Packaging")):
    iid = iid_of(d)
    pkg.append(dict(
        date=ddmm(col(d, "Date and Time - Date", "Date -")),
        cats=str(col(d, "General Product Categories") or "").strip(),
        guard=str(col(d, "Full name") or "").strip(),
        val=PKG_VALUES_BY_ID.get(iid, 0.0),
    ))
pkg.sort(key=lambda x: -x["val"])

# ---------- metrics ----------
stores = sorted(set(i["store"] for i in inc))
N = len(inc)
interv = sum(1 for i in inc if i["cat"] == "Intervention")
deter = sum(1 for i in inc if i["cat"] == "Deterrent")
other = N - interv - deter
vt = sum(1 for i in inc if i["vt"].lower() == "yes")
INC_REC = sum(i["rec"] for i in inc)
PKG_TOTAL = sum(p["val"] for p in pkg)
PKG_VALUED = sum(1 for p in pkg if p["val"])
TOTAL_REC = INC_REC + PKG_TOTAL

def m(v): return f"${v:,.2f}"
logo = f'<img src="data:image/png;base64,{base64.b64encode(open(LOGO,"rb").read()).decode()}" style="height:34px;opacity:.9"/>' if os.path.exists(LOGO) else ""

def kpi(l, v, c="#1f2933", s=""):
    sub = f'<div style="font-size:8.5px;color:{SUB}">{s}</div>' if s else ""
    return (f'<div style="flex:1"><div style="font-size:8.5px;color:{SUB};text-transform:uppercase;letter-spacing:.6px">{l}</div>'
            f'<div style="font-size:21px;font-weight:700;color:{c};margin-top:2px">{v}</div>{sub}</div>')

TH = f'background:{CARD};color:{SUB};font-size:8.5px;text-transform:uppercase;letter-spacing:.5px;padding:6px 8px'
BASE = ("*{margin:0;box-sizing:border-box;font-family:-apple-system,Segoe UI,Arial,sans-serif}"
        "body{background:#fff;color:#1f2933;padding:24px 28px}table{width:100%;border-collapse:collapse;font-size:11px}"
        f"tbody tr{{border-bottom:1px solid {LINE}}}h2{{color:{ACC};font-size:13px;letter-spacing:.6px;margin:0 0 6px}}"
        f".lead{{color:{SUB};font-size:11px;line-height:1.5}}")

def inc_rows():
    o = ""
    for i in inc:
        catc = GREEN if i["cat"] == "Intervention" else (ACC if i["cat"] == "Deterrent" else SUB)
        rc = m(i["rec"]) if i["rec"] else "—"
        o += (f'<tr><td style="padding:4px 8px">{i["date"]}, {i["time"]}</td>'
              f'<td style="text-align:center;color:{catc};font-weight:700">{i["cat"]}</td>'
              f'<td style="padding:4px 8px;color:{SUB}">{i["cause"]}</td>'
              f'<td style="text-align:center;color:{RED if i["vt"].lower()=="yes" else SUB};font-weight:{700 if i["vt"].lower()=="yes" else 400}">{i["vt"]}</td>'
              f'<td style="text-align:right;color:{GREEN if i["rec"] else SUB};font-weight:{600 if i["rec"] else 400}">{rc}</td></tr>')
    o += (f'<tr style="border-top:2px solid {GREEN}"><td colspan="4" style="padding:6px 8px;font-weight:700">RECOVERED AT INCIDENTS</td>'
          f'<td style="text-align:right;font-weight:700;color:{GREEN}">{m(INC_REC)}</td></tr>')
    return o

def pkg_rows():
    o = ""
    for p in pkg:
        rc = m(p["val"]) if p["val"] else "—"
        o += (f'<tr><td style="padding:4px 8px;color:{SUB}">{p["date"]}</td>'
              f'<td style="padding:4px 8px;color:{SUB};font-size:9.5px">{p["cats"]}</td>'
              f'<td style="text-align:right;color:{GREEN if p["val"] else SUB};font-weight:{600 if p["val"] else 400}">{rc}</td></tr>')
    o += (f'<tr style="border-top:2px solid {GREEN}"><td colspan="2" style="padding:6px 8px;font-weight:700">TOTAL</td>'
          f'<td style="text-align:right;font-weight:700;color:{GREEN}">{m(PKG_TOTAL)}</td></tr>')
    return o

HTML = f"""<html><head><meta charset='utf-8'><style>{BASE}</style></head><body>
<div style="display:flex;justify-content:space-between;align-items:flex-start;border-bottom:2px solid {ACC};padding-bottom:10px">
  <div><div style="font-size:21px;font-weight:800">THE WAREHOUSE — WEEKLY COVERT LP REPORT</div>
  <div style="color:{SUB};font-size:12px;margin-top:2px">The Warehouse Takanini · SLS RTC Limited · {PERIOD}</div></div>
  {logo}</div>
<div style="display:flex;gap:6px;background:{CARD};border:1px solid {LINE};border-radius:8px;padding:12px 16px;margin:14px 0">
  {kpi("Store", "1", ACC, "Takanini")}{kpi("Incidents", str(N))}{kpi("Interventions", str(interv), GREEN)}{kpi("Deterrents", str(deter), ACC)}{kpi("Verbal Trespasses", str(vt), RED if vt else "#1f2933")}{kpi("Recovered", m(INC_REC), GREEN, f"{sum(1 for i in inc if i['rec'])} incidents")}{kpi("Recovered Pkg", m(PKG_TOTAL), GREEN, f"{PKG_VALUED} finds")}{kpi("Total Recovered", m(TOTAL_REC), ACC)}</div>
<p class="lead">SLS RTC continued covert loss-prevention at <b>The Warehouse Takanini</b> this week. The team actioned <b>{N} incidents</b>
({interv} interventions, {deter} deterrents), issued <b>{vt} verbal trespass notices</b>, recovered <b>{m(INC_REC)}</b> of product at incidents and a
further <b>{m(PKG_TOTAL)}</b> of product packaging as theft evidence — <b>{m(TOTAL_REC)}</b> total — all logged for submission into Auror, the cumulative
offending record that builds the threshold for Police escalation.</p>
<div style="display:flex;gap:18px;margin-top:14px">
  <div style="flex:1.9"><h2>INCIDENT LOG</h2>
  <table><thead><tr><th style="{TH};text-align:left">Date/Time</th><th style="{TH};text-align:center">Category</th><th style="{TH};text-align:left">Cause</th><th style="{TH};text-align:center">Verbal Trespass</th><th style="{TH};text-align:right">Recovered</th></tr></thead><tbody>{inc_rows()}</tbody></table></div>
  <div style="flex:1.1"><h2 style="color:{GREEN}">RECOVERED PACKAGING — {m(PKG_TOTAL)}</h2>
  <table><thead><tr><th style="{TH};text-align:left">Date</th><th style="{TH};text-align:left">Categories</th><th style="{TH};text-align:right">Receipt Value</th></tr></thead><tbody>{pkg_rows()}</tbody></table>
  <div style="font-size:9px;color:{SUB};margin-top:5px">{len(pkg)} packaging finds logged; receipt values confirmed on {PKG_VALUED}.</div></div></div>
<div style="margin-top:16px;background:{CARD};border-left:3px solid {ACC};border-radius:4px;padding:10px 13px;font-size:9.5px;color:{SUB};line-height:1.55">
  <b style="color:{ACC}">Week 2.</b> "Recovered" is the value of product recovered at the incident; "recovered packaging" is the receipt value of empty product packaging located in-aisle —
  quantified theft evidence. Both are logged to Auror. Interventions stop theft in progress; deterrents are surveillance-driven prevention (suspects identified the covert team and exited).
  Source: SLS RTC covert LP reporting (Auror-linked incidents).</div>
</body></html>"""

def render(html, out):
    tmp = os.path.join(WEEK, "_tmp.html"); open(tmp, "w").write(html)
    subprocess.run([CHROME, "--headless", "--disable-gpu", "--no-pdf-header-footer",
                    f"--print-to-pdf={out}", "--no-margins", tmp], capture_output=True, timeout=90)
    os.remove(tmp); print("Wrote", out)

OUT1 = os.path.join(WEEK, f"{PREFIX}_LP_Report.pdf")
render(HTML, OUT1)

# ---------- email ----------
from docx import Document
from docx.shared import Pt, RGBColor
doc = Document(); doc.styles["Normal"].font.name = "Calibri"; doc.styles["Normal"].font.size = Pt(11)
def P(t, b=False, sz=11, c=None, sa=6):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(sa)
    r = p.add_run(t); r.bold = b; r.font.size = Pt(sz)
    if c: r.font.color.rgb = RGBColor(*c)
def Bul(t):
    p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(3); p.add_run(t)
P("To:        The Warehouse — Loss Prevention", sa=0)
P("From:    SLS RTC Limited", sa=0)
P(f"Subject: The Warehouse Takanini — Covert LP Weekly Summary ({PERIOD})", b=True, c=(0x1f,0x3a,0x5f))
P("Hi team,")
P("Please find this week's covert loss-prevention summary for The Warehouse Takanini, with the supporting report attached for "
  "submission into Auror — the platform where the cumulative record of offending is built, breaking the repeat-offender cycle and "
  "establishing the evidence threshold required for Police escalation.")
P("Week in brief", b=True, c=(0x1f,0x3a,0x5f), sa=4)
Bul(f"{N} incidents actioned ({interv} interventions, {deter} deterrents), with {vt} verbal trespass notices issued.")
Bul(f"{m(INC_REC)} of product recovered at incidents, plus {m(PKG_TOTAL)} of recovered packaging logged as theft evidence — {m(TOTAL_REC)} total for the week.")
Bul("Deterrents this week were surveillance-driven — suspects identified the covert team and exited before offending.")
P("Attached", b=True, c=(0x1f,0x3a,0x5f), sa=4)
Bul("Weekly LP Report — incident activity and recovered packaging for the week.")
P("Happy to jump on a quick call to walk through the week and the plan ahead.")
P("Many thanks,", sa=2)
P("Ahmad Duais — Director, SLS RTC Limited", sa=0)
P(f"Attachment: {os.path.basename(OUT1)}", sz=9, c=(0x5b,0x64,0x72))
OUT2 = os.path.join(WEEK, f"{PREFIX}_Email_Summary.docx")
doc.save(OUT2); print("Wrote", OUT2)
print(f"\nSUMMARY: incidents={N} interv={interv} deter={deter} vt={vt} inc_rec={m(INC_REC)} pkg={m(PKG_TOTAL)} ({PKG_VALUED}/{len(pkg)}) total={m(TOTAL_REC)}")
