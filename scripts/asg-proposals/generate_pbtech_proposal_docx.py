#!/usr/bin/env python3
"""
ASG Security — PB Tech Integrated Security Proposal, editable Word (.docx) edition.
Content mirrors generate_pbtech_proposal.py (the PDF) verbatim — keep both in sync.
Output: ~/Desktop/ASG/Sales & Marketing/Proposals/PBTech_Integrated_Security_Proposal.docx
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Mm, Pt, RGBColor

MIDNIGHT = RGBColor(0x05, 0x0A, 0x14)
SIGNAL = RGBColor(0x25, 0x63, 0xEB)
SLATE = RGBColor(0x33, 0x41, 0x55)
GRAY_400 = RGBColor(0x94, 0xA3, 0xB8)
GRAY_500 = RGBColor(0x64, 0x74, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BODY_FONT = 'Arial'
MONO_FONT = 'Consolas'

OUT = Path.home() / 'Desktop/ASG/Sales & Marketing/Proposals/PBTech_Integrated_Security_Proposal.docx'


def shade(cell, hexfill):
    cell._tc.get_or_add_tcPr().append(
        parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:fill="{hexfill}"/>'))


def set_borders(table, edges, size=4, color='E2E8F0'):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in edges:
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), str(size))
        el.set(qn('w:color'), color)
        borders.append(el)
    tbl_pr.append(borders)


def add_runs(p, text, base_color=SLATE, size=10, font=BODY_FONT):
    """Render <b>...</b> markup into Word runs."""
    text = text.replace('&amp;', '&').replace('&nbsp;', ' ')
    for chunk in re.split(r'(<b>.*?</b>)', text, flags=re.S):
        if not chunk:
            continue
        bold = chunk.startswith('<b>')
        seg = re.sub(r'</?b>', '', chunk)
        for j, part in enumerate(seg.split('<br/>')):
            if j:
                p.add_run().add_break()
            r = p.add_run(part)
            r.bold = bold
            r.font.name = font
            r.font.size = Pt(size)
            r.font.color.rgb = MIDNIGHT if bold else base_color
    return p


def body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.12
    return add_runs(p, text)


def bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.12
    return add_runs(p, text)


def heading(doc, num, text, space_before=10):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(num)
    r.bold = True
    r.font.name = BODY_FONT
    r.font.size = Pt(13.5)
    r.font.color.rgb = SIGNAL
    r2 = p.add_run('  ' + text)
    r2.bold = True
    r2.font.name = BODY_FONT
    r2.font.size = Pt(13.5)
    r2.font.color.rgb = MIDNIGHT
    return p


def subheading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.bold = True
    r.font.name = BODY_FONT
    r.font.size = Pt(10.5)
    r.font.color.rgb = MIDNIGHT
    return p


def title_block(doc):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    cell = t.cell(0, 0)
    cell.width = Mm(172)
    shade(cell, '050A14')
    set_borders(t, ['top', 'left', 'bottom', 'right'], size=4, color='050A14')

    p1 = cell.paragraphs[0]
    p1.paragraph_format.space_before = Pt(14)
    p1.paragraph_format.space_after = Pt(16)
    r = p1.add_run('ASG')
    r.bold = True; r.font.name = BODY_FONT; r.font.size = Pt(14); r.font.color.rgb = WHITE
    r = p1.add_run('  |  ')
    r.font.name = BODY_FONT; r.font.size = Pt(12); r.font.color.rgb = GRAY_500
    r = p1.add_run('S E C U R I T Y')
    r.font.name = BODY_FONT; r.font.size = Pt(10); r.font.color.rgb = GRAY_400
    p1.add_run('\t')
    r = p1.add_run('Complete Security. One Partner.')
    r.font.name = BODY_FONT; r.font.size = Pt(8.5); r.font.color.rgb = GRAY_400
    p1.paragraph_format.tab_stops.add_tab_stop(Mm(165), WD_TAB_ALIGNMENT.RIGHT)

    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(4)
    r = p2.add_run('I N T E G R A T E D   S E C U R I T Y   P R O P O S A L')
    r.bold = True; r.font.name = BODY_FONT; r.font.size = Pt(8.5); r.font.color.rgb = SIGNAL

    p3 = cell.add_paragraph()
    p3.paragraph_format.space_after = Pt(10)
    r = p3.add_run('PB Tech — nationwide security operations, one partner')
    r.bold = True; r.font.name = BODY_FONT; r.font.size = Pt(22); r.font.color.rgb = WHITE

    p4 = cell.add_paragraph()
    p4.paragraph_format.space_after = Pt(2)
    r = p4.add_run('Prepared for PB Technologies Ltd    ·    October 2026    ·    Commercial in confidence')
    r.font.name = BODY_FONT; r.font.size = Pt(9.5); r.font.color.rgb = GRAY_400
    p5 = cell.add_paragraph()
    p5.paragraph_format.space_after = Pt(14)
    r = p5.add_run('Absolute Security Group Ltd  &  SLS RTC Ltd')
    r.font.name = BODY_FONT; r.font.size = Pt(9.5); r.font.color.rgb = GRAY_400


def stat_strip(doc):
    stats = [
        ('100+', 'retail sites under covert LP'),
        ('$1M+', 'verified recoveries since Jul 2025'),
        ('48 hr', 'covert mobilisation, base regions'),
        ('2006', 'SLS RTC in retail LP since'),
    ]
    t = doc.add_table(rows=2, cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for row in t.rows:
        for c in row.cells:
            c.width = Mm(43)
    set_borders(t, ['top', 'bottom'], size=6, color='E2E8F0')
    for i, (num, label) in enumerate(stats):
        pn = t.cell(0, i).paragraphs[0]
        pn.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pn.paragraph_format.space_before = Pt(6)
        r = pn.add_run(num)
        r.bold = True; r.font.name = MONO_FONT; r.font.size = Pt(15); r.font.color.rgb = MIDNIGHT
        pl = t.cell(1, i).paragraphs[0]
        pl.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pl.paragraph_format.space_after = Pt(6)
        r = pl.add_run(label)
        r.font.name = BODY_FONT; r.font.size = Pt(8); r.font.color.rgb = GRAY_500


def impl_table(doc):
    rows = [
        ['Phase', 'Timing', 'Scope'],
        ['0 — Survey & design', 'Weeks 1–2',
         'Site surveys; risk tiering from PB Tech loss data (under NDA); burglary/ram-raid hardening '
         'assessment; SOPs and duress protocol; privacy scoping across CCTV, crime-intelligence and FRT'],
        ['1 — Static & monitoring', 'Weeks 3–6',
         'Guarding live at priority stores; monitoring transition aligned to incumbent notice periods, with '
         'parallel running; reporting live from first shift'],
        ['2 — Covert LP', 'Weeks 5–10',
         'Covert rotation live nationally from the four regional bases; baseline accumulating; first tier '
         're-weighting at week 8'],
        ['3 — FRT (gated)', 'Week 8 on',
         'Necessity-and-proportionality assessment completed with PB Tech; pilot proceeds only on PB Tech '
         'go/no-go, in 2–3 stores with a measured false-positive gate'],
    ]
    t = doc.add_table(rows=len(rows), cols=3)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Mm(42), Mm(25), Mm(108)]
    for row in t.rows:
        for c, w in zip(row.cells, widths):
            c.width = w
    for j, cell in enumerate(t.rows[0].cells):
        shade(cell, '050A14')
        p = cell.paragraphs[0]
        r = p.add_run(rows[0][j])
        r.bold = True; r.font.name = BODY_FONT; r.font.size = Pt(8.5); r.font.color.rgb = WHITE
    for i, row in enumerate(rows[1:], start=1):
        if i % 2 == 0:
            for cell in t.rows[i].cells:
                shade(cell, 'F1F5F9')
        for j, val in enumerate(row):
            p = t.rows[i].cells[j].paragraphs[0]
            r = p.add_run(val)
            r.bold = (j == 0)
            r.font.name = BODY_FONT
            r.font.size = Pt(8.5)
            r.font.color.rgb = MIDNIGHT if j == 0 else SLATE


def footer(doc):
    sec = doc.sections[0]
    p = sec.footer.paragraphs[0]
    p.paragraph_format.tab_stops.add_tab_stop(Mm(170), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run('ASG Security  ·  Commercial in confidence\t')
    r.font.name = BODY_FONT; r.font.size = Pt(7.5); r.font.color.rgb = GRAY_500
    rp = p.add_run('Page ')
    rp.font.name = BODY_FONT; rp.font.size = Pt(7.5); rp.font.color.rgb = GRAY_500
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    p._p.append(fld)


def main():
    doc = Document()
    doc.core_properties.title = 'PB Tech — Integrated Security Proposal'
    doc.core_properties.author = 'ASG Security'
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(19)
    sec.top_margin, sec.bottom_margin = Mm(20), Mm(18)
    style = doc.styles['Normal']
    style.font.name = BODY_FONT
    style.font.size = Pt(10)
    style.font.color.rgb = SLATE

    title_block(doc)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    heading(doc, '1', 'Executive summary', space_before=4)
    body(doc, 'Security for a nationwide retail network is usually bought as three separate services — '
               'guarding from one vendor, alarm monitoring from another, loss prevention from a third — none of '
               'which share data. The gaps between vendors are where loss happens. ASG Group proposes to run all of '
               'it for PB Tech as one partnership: <b>one vendor, one SLA, one weekly report, every store.</b>')
    body(doc, 'The risk profile justifies it. PB Tech’s superstores and service centres carry dense '
               'concentrations of high-value, high-resale, serialised stock — a profile that draws organised '
               'repeat offenders working across branches, distraction theft at counters, refund and warranty '
               'fraud, and targeted after-hours attack on stores and warehouses. The partnership is built on '
               'five layers:')
    bullet(doc, '<b>Risk-tiered static guarding</b> at the stores where measured loss justifies a visible presence.')
    bullet(doc, '<b>Covert loss prevention</b> — plain-clothes operatives who engage offenders where it is '
                'safe and lawful, recover stock, and document everything.')
    bullet(doc, '<b>CCTV and alarm monitoring</b>, 24/7, with video-verified response wherever live footage '
                'is available.')
    bullet(doc, '<b>Facial recognition (optional, privacy-gated)</b> — watchlist alerting for serious repeat '
                'offenders, taken forward only if a joint assessment under the Biometric Processing Privacy '
                'Code 2025 supports it.')
    bullet(doc, '<b>Data-led coverage</b> — guard-hours re-weighted every week from your own incident and '
                'recovery data, so cover sits where loss is occurring.')
    body(doc, '<b>Who we are.</b> ASG Group is New Zealand-owned. Its covert loss-prevention arm, '
               '<b>SLS RTC Ltd</b>, has protected NZ retail continuously since 2006 and joined the group in '
               '2025, alongside <b>Absolute Security Group Ltd</b> (uniformed guarding, precinct patrol, CCTV '
               'monitoring and electronic security). 180+ security professionals on active rosters '
               'group-wide; offices in Auckland, Wellington and Christchurch; covert teams based in Auckland, '
               'Hamilton, Wellington and Christchurch, working every week inside national retail networks '
               'across grocery, hardware and general merchandise. Client data is strictly siloed — never '
               'cross-shared — and a dedicated PB Tech team with contractual information barriers is part of '
               'this offer.')
    stat_strip(doc)

    heading(doc, '2', 'Capability — what the group runs today')
    body(doc, '<b>Covert retail loss prevention (SLS RTC).</b> Two-operative minimum teams as standing '
               'policy, rotating across national retail networks. Operatives are COA-certified, SOP-trained '
               'and crime-intelligence-platform trained; they engage and recover under a documented '
               'use-of-force and citizen’s-arrest policy, administer trespass warnings as the client’s '
               'authorised agents under the Trespass Act 1980, and build structured evidence packs for '
               'Police. Recoveries are reconciled weekly with each client — more than $1 million in '
               'verified recoveries across the covert portfolio since July 2025 — and every incident is '
               'graded on our 1–5 severity scale (1 = compliant through 5 = physical/weapons), so you see '
               'the safety picture your staff face, not just the recovery number.')
    body(doc, '<b>Guarding, precinct and buildings (ASG).</b> Retail, corporate, venue and precinct work, '
               'including a multi-year city-centre precinct safety patrol contract — thousands of patrol '
               'routes and business-assistance calls completed, with zero use of force recorded across every '
               'aggression incident in the life of the contract. On our longest-running commercial building '
               'contract, over half of all logged activity is preventative — issues found and fixed before '
               'they become incidents. Surge capability is proven: simultaneous crowd-control activation '
               'across multiple stores in five regions, stood up inside two weeks.')
    body(doc, '<b>CCTV, alarm monitoring and electronic security.</b> We operate on the Nx Witness video '
               'management platform with in-house alert-triage engineering, so operators work a ranked, '
               'verified queue rather than a raw alarm flood — and camera-health monitoring is standing '
               'practice: where a camera sits within our monitored estate, an outage is logged, reported and '
               'escalated with the discipline of an incident. The hardware side is in-house too: EAS gates '
               '(RF and AM), display-stock protection for high-value devices, tags, labels and detachers — '
               'supplied, installed and serviced nationwide — plus the Mosquito anti-loitering system, alarm '
               'response and a mobile patrol fleet.')
    body(doc, '<b>People and reporting.</b> Officers are trained through our in-house academy, and '
               'turnover across our permanent security workforce ran under 10% in the last 12 months. We do '
               'not subcontract outside the group — every officer is directly employed by an ASG Group '
               'company, every hour timeclock-verified; the head-contractor structure is set out in Section '
               '6. Daily, weekly and monthly reporting is standard on every contract, and — where authorised '
               'by the client and consistent with its privacy policy — we integrate with retail '
               'crime-intelligence platforms such as Auror.')

    heading(doc, '3', 'Proposed configuration for PB Tech')
    body(doc, '<b>3.1  Static guarding — risk-tiered, not blanket.</b> We do not propose a guard in '
               'every store. Phase 0 tiering, built from your loss data and our site surveys, sets which sites '
               'carry trading-hour presence, which carry peak-window presence (weekends, launch days, '
               'promotional periods), and which are covered by covert rotation and monitoring. The protocol for '
               'when your staff face aggression — including duress-alarm response — is agreed in Phase 0, '
               'alongside the overlapping-duties consultation the Health and Safety at Work Act requires '
               'between both PCBUs. Electronics-specific training modules (distraction theft, grab-and-run on '
               'display stock, refund and warranty fraud, service-centre custody of customer devices) are '
               'built with your input in Phase 0.')
    body(doc, '<b>3.2  Covert loss prevention — intercept and recover.</b> Rotating covert teams '
               'concentrate on the highest-loss stores, with depth set by measured return (Section 4). '
               'Operatives engage and recover where it is safe and lawful — our staff hold no detention powers '
               'beyond any citizen’s, and are trained to de-escalate, preserve evidence and refer to Police '
               'rather than confront — administer trespass warnings under written delegation, and build '
               'offender files. Because one team works many branches, cross-store patterns — the same group '
               'hitting your Albany branch on Tuesday and St Lukes on Thursday — surface in days, not months. '
               'Where loss patterns point inside, the same team runs covert investigations and serialised-stock '
               'tracing.')
    body(doc, '<b>3.3  CCTV and alarm monitoring — verified response.</b> 24/7 monitoring of alarm '
               'activations. Events are video-verified where live footage is available; otherwise the agreed '
               'alarm-response protocol applies — verification never delays response. Patrol or keyholder '
               'response to verified events. We integrate with your existing camera estate where it is fit for '
               'purpose and specify upgrades only where coverage gaps create risk. Because after-hours attack '
               'is a defining risk for stores and warehouses full of resaleable stock, Phase 0 includes a '
               'site-by-site burglary and ram-raid hardening assessment, and after-hours response targets are '
               'set per location class in the SLA.')
    body(doc, '<b>3.4  Facial recognition — staged, honest, gated.</b> FRT in New Zealand retail is '
               'governed by the Biometric Processing Privacy Code 2025; done wrong it is a regulatory and '
               'reputational liability, and the evidence from New Zealand’s retail trials to date is '
               'qualified — which is why we pilot and measure rather than promise. How it works, stated '
               'plainly: every entering face is checked against a watchlist; non-matches are deleted '
               'immediately; only enrolled individuals are retained. Enrolment is limited to serious repeat '
               'offenders — formally trespassed, identity verified, approved person-by-person by PB Tech, no '
               'minors or vulnerable persons, entries expiring with the two-year trespass warning and reviewed '
               'on a set cycle. Every alert is verified by a human operator before any floor response; no '
               'action is ever taken on an automated match alone. Controls include purpose limitation, defined '
               'retention, role-based access with audit logs, accuracy testing before go-live, clear entry '
               'signage, and incident and breach procedures. PB Tech retains the go/no-go decision on the '
               'necessity-and-proportionality assessment — we prepare it with you and supply the operational '
               'evidence, with each party responsible for its own obligations under the Privacy Act and the '
               'Code. If the assessment does not support FRT, we do not run it; the rest of this proposal '
               'stands on its own.')

    heading(doc, '4', 'Data-led coverage — hours follow the loss')
    body(doc, 'Most guarding contracts fix hours at tender and never move them. Ours re-weights cover every '
               'week from your own incident and recovery data.')
    body(doc, '<b>Baseline and tiering.</b> Weeks 1–8 build a per-store baseline: incidents, recoveries, '
               'offender counts, trespass register, alarm activations and severity spread. Stores are tiered '
               'from the data, not by intuition.')
    body(doc, '<b>Return-weighted coverage.</b> Covert days and depth per store are set after the '
               'baseline, weighing recovery value, incident severity, loss trend, stock profile and safety '
               'risk — with a breadth rule guaranteeing no store is ever left uncovered. The weighting is '
               'recalculated weekly, so when offenders shift to a quieter branch, the hours follow them '
               'within a week. Proposed hours per store sit in the priced plan, not in this document.')
    body(doc, '<b>Evidence pipeline.</b> Every documented incident feeds offender profiles and structured '
               'evidence packs for Police — subject always to Police evidential requirements and discretion. '
               'The trespass register feeds the FRT watchlist only under the enrolment standard in 3.4.')
    subheading(doc, 'Weekly reporting package — one pack, every Tuesday')
    bullet(doc, '<b>Deployment &amp; Recovery Tracker (Excel)</b> — hours, recoveries, incidents, offenders, '
                'deterrents and trespasses per store; cost against budget, week / period / year-to-date.')
    bullet(doc, '<b>National summary</b> — store, region and national totals with week-on-week movements '
                'explained, not just reported. Billed hours reconcile to timeclock records, recoveries are '
                'valued at verified retail value, and totals reconcile exactly across every level of the pack.')
    bullet(doc, '<b>Safety report</b> — severity spread on the 1–5 scale, serious (3+) incidents narrated, '
                'Police referrals tracked.')
    body(doc, '<b>Governance cadence.</b> Weekly report → monthly operations review with your LP lead → '
               'quarterly strategy review where tiering, static placement and the FRT gate are re-set from '
               'the data.')

    heading(doc, '5', 'Implementation')
    impl_table(doc)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

    heading(doc, '6', 'Accountability, compliance and commercial approach')
    body(doc, '<b>SLA.</b> Weekly report every Tuesday; monthly operations review; quarterly business '
               'reviews with the data on the table. Shift-fill guarantee and monitored-event response targets '
               'by location class, backed by service credits — the numbers are tabled with the priced plan.')
    body(doc, '<b>Contracting.</b> Absolute Security Group Ltd contracts as head contractor; SLS RTC Ltd '
               'delivers covert LP under it as a disclosed subcontractor with flow-down obligations — one '
               'counterparty, one invoice.')
    body(doc, '<b>Licensing, vetting and insurance.</b> Both entities are licensed under the Private '
               'Security Personnel and Private Investigators Act 2010 in the classes each service requires; '
               'every operative holds the required current licence or Certificate of Approval, verified at '
               'induction and re-checked during the term. Pre-employment vetting includes Police checks. '
               'Public liability and statutory liability cover is held by both entities — certificates of '
               'currency and limits are provided with the priced plan.')
    body(doc, '<b>Your data.</b> PB Tech’s incident, offender and loss data lives in an access-controlled '
               'store siloed from every other client, with role-based access, return or destruction on exit, '
               'and breach notification agreed in the SLA.')
    body(doc, '<b>Commercial.</b> A single national agreement, one consolidated monthly invoice, and '
               'cost-versus-budget visible in every weekly tracker — you should never have to ask what you '
               'spent or what it returned. Billing is on timeclock-verified actual hours only. Annual rate '
               'adjustments are capped at the lower of 2.5% and the annual movement in Stats NZ’s Labour '
               'Cost Index, with statutory minimum-wage changes passed through at demonstrated cost only — '
               'cost certainty for the term, not mid-term renegotiation.')

    heading(doc, '7', 'The offer — how we start')
    body(doc, '<b>Start free: the Loss-Map Audit.</b> A walkthrough of three representative sites (one '
               'superstore, one service centre, one distribution point) and a loss-data workshop with your LP '
               'lead — no cost, no obligation. We are on site within five working days of your confirmation '
               'and NDA execution, and return a store-by-store coverage plan, risk tiering and an indicative '
               'return scenario (assumptions stated) within ten working days of the workshop and receipt of '
               'your loss data. An indicative rate card per service line is available on request, independent '
               'of the audit.')
    body(doc, '<b>Pilot without risk.</b> A four-week covert LP pilot in your highest-loss stores, live '
               'within 48 hours of pilot-scope sign-off and store-access confirmation in our four covert base '
               'regions (other regions within five working days). Our pilot guarantee: if verified recoveries '
               'and documented interventions across the four weeks don’t exceed the pilot’s invoiced cost — '
               'measured against a baseline agreed before day one, on the weekly tracker both parties sign '
               'off — you don’t pay. The pilot runs ahead of Phase 2: its stores, data and weeks roll '
               'straight into the per-store baseline, and pilot teams transition into the national rotation '
               'without a restart.')
    body(doc, '<b>You keep the artefacts.</b> Everything the Loss-Map Audit produces — the walkthrough '
               'reports, draft risk tiering and indicative return scenario — is yours to keep, whatever you '
               'decide. Phase 0 then validates that tiering with full site surveys and adds the SOPs, '
               'hardening assessments and privacy pack.')
    body(doc, '<b>Timing.</b> On a Phase 0 start by Friday 13 November 2026, guarding and monitoring are '
               'live at priority stores by early December and run through the peak, covert pilot teams can be '
               'in stores ahead of Black Friday on a prompt decision, and the full national rotation '
               'completes by late January. The proposal remains open for 60 days; after mid-November, a '
               'late-January full start remains available with peak coverage resting on the pilot.')
    body(doc, '<b>Ahmad Duais</b> — Chief Executive, Absolute Security Group<br/>'
               'ahmad@absolutesecuritygroup.co.nz  ·  absolutesecuritygroup.co.nz  ·  slsrtc.co.nz')

    footer(doc)
    doc.save(str(OUT))
    print(f'OK {OUT}')


if __name__ == '__main__':
    main()
