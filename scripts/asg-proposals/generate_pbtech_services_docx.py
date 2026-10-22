#!/usr/bin/env python3
"""
ASG Security | PB Tech 2.0: Security Services Overview (editable Word).
Companion to PBTech_Integrated_Security_Proposal.docx (the v1 proposal).
Modelled on the SLS RTC services-deck structure (challenge, layers, operating
model, systems, reporting, compliance, results, people, engagement, contact).
No em-dashes anywhere (owner preference). No rates. Verified claims + the
owner-approved overrides (350+, under-5% turnover) only.
Output: ~/Desktop/ASG/Sales & Marketing/Proposals/PBTech_Security_Services_Overview_2.0.docx
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Mm, Pt, RGBColor

MIDNIGHT = RGBColor(0x05, 0x0A, 0x14)
SIGNAL = RGBColor(0x25, 0x63, 0xEB)
SLATE = RGBColor(0x33, 0x41, 0x55)
GRAY_400 = RGBColor(0x94, 0xA3, 0xB8)
GRAY_500 = RGBColor(0x64, 0x74, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BODY_FONT = 'Arial'
MONO_FONT = 'Consolas'

OUT = Path.home() / 'Desktop/ASG/Sales & Marketing/Proposals/PBTech_Security_Services_Overview_2.0.docx'
TEAM_PHOTO = Path(__file__).parent / 'assets/asg_leadership_team.png'


def shade(cell, hexfill):
    cell._tc.get_or_add_tcPr().append(
        parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:fill="{hexfill}"/>'))


def set_borders(table, edges, size=4, color='E2E8F0'):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in edges:
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), str(size))
        el.set(qn('w:color'), color)
        borders.append(el)
    tbl_pr.append(borders)


def add_runs(p, text, base_color=SLATE, size=10, font=BODY_FONT):
    text = text.replace('&amp;', '&')
    for chunk in re.split(r'(<b>.*?</b>)', text, flags=re.S):
        if not chunk:
            continue
        bold = chunk.startswith('<b>')
        seg = re.sub(r'</?b>', '', chunk)
        r = p.add_run(seg)
        r.bold = bold
        r.font.name = font
        r.font.size = Pt(size)
        r.font.color.rgb = MIDNIGHT if bold else base_color
    return p


def body(doc, text, size=10):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.12
    return add_runs(p, text, size=size)


def bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.12
    return add_runs(p, text)


def overline(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(' '.join(text.upper()))
    r.bold = True
    r.font.name = BODY_FONT
    r.font.size = Pt(8)
    r.font.color.rgb = SIGNAL
    return p


def heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.bold = True
    r.font.name = BODY_FONT
    r.font.size = Pt(15)
    r.font.color.rgb = MIDNIGHT
    return p


def cell_text(cell, runs, align=None):
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    for text, kw in runs:
        r = p.add_run(text)
        r.bold = kw.get('bold', False)
        r.font.name = kw.get('font', BODY_FONT)
        r.font.size = Pt(kw.get('size', 9))
        r.font.color.rgb = kw.get('color', SLATE)
    return p


def cover(doc):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    cell = t.cell(0, 0)
    cell.width = Mm(172)
    shade(cell, '050A14')
    set_borders(t, ['top', 'left', 'bottom', 'right'], size=4, color='050A14')

    p1 = cell.paragraphs[0]
    p1.paragraph_format.space_before = Pt(14)
    p1.paragraph_format.space_after = Pt(16)
    r = p1.add_run('ASG')
    r.bold = True; r.font.name = BODY_FONT; r.font.size = Pt(14); r.font.color.rgb = WHITE
    r = p1.add_run('  |  ')
    r.font.name = BODY_FONT; r.font.size = Pt(12); r.font.color.rgb = GRAY_500
    r = p1.add_run('S E C U R I T Y')
    r.font.name = BODY_FONT; r.font.size = Pt(10); r.font.color.rgb = GRAY_400
    p1.add_run('\t')
    r = p1.add_run('Complete Security. One Partner.')
    r.font.name = BODY_FONT; r.font.size = Pt(8.5); r.font.color.rgb = GRAY_400
    p1.paragraph_format.tab_stops.add_tab_stop(Mm(165), WD_TAB_ALIGNMENT.RIGHT)

    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(4)
    r = p2.add_run('S E C U R I T Y   S E R V I C E S   O V E R V I E W')
    r.bold = True; r.font.name = BODY_FONT; r.font.size = Pt(8.5); r.font.color.rgb = SIGNAL

    p3 = cell.add_paragraph()
    p3.paragraph_format.space_after = Pt(6)
    r = p3.add_run('Integrated Security and Loss Prevention')
    r.bold = True; r.font.name = BODY_FONT; r.font.size = Pt(22); r.font.color.rgb = WHITE

    p4 = cell.add_paragraph()
    p4.paragraph_format.space_after = Pt(10)
    r = p4.add_run('Covert loss prevention, uniformed guarding, electronic security, monitoring and '
                   'intelligence-led reporting for one of New Zealand’s largest technology retailers.')
    r.font.name = BODY_FONT; r.font.size = Pt(10.5); r.font.color.rgb = GRAY_400

    p5 = cell.add_paragraph()
    p5.paragraph_format.space_after = Pt(2)
    r = p5.add_run('Prepared for PB Technologies Ltd    ·    October 2026    ·    Commercial in confidence')
    r.font.name = BODY_FONT; r.font.size = Pt(9.5); r.font.color.rgb = GRAY_400
    p6 = cell.add_paragraph()
    p6.paragraph_format.space_after = Pt(14)
    r = p6.add_run('Absolute Security Group  &  SLS Retail Theft Control')
    r.font.name = BODY_FONT; r.font.size = Pt(9.5); r.font.color.rgb = GRAY_400


def stat_strip(doc):
    stats = [
        ('350+', 'security professionals group-wide'),
        ('$1M+', 'verified recoveries since Jul 2025'),
        ('100+', 'retail sites under covert LP'),
        ('2006', 'SLS RTC protecting NZ retail since'),
    ]
    t = doc.add_table(rows=2, cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for row in t.rows:
        for c in row.cells:
            c.width = Mm(43)
    set_borders(t, ['top', 'bottom'], size=6, color='E2E8F0')
    for i, (num, label) in enumerate(stats):
        pn = t.cell(0, i).paragraphs[0]
        pn.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pn.paragraph_format.space_before = Pt(6)
        r = pn.add_run(num)
        r.bold = True; r.font.name = MONO_FONT; r.font.size = Pt(15); r.font.color.rgb = MIDNIGHT
        pl = t.cell(1, i).paragraphs[0]
        pl.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pl.paragraph_format.space_after = Pt(6)
        r = pl.add_run(label)
        r.font.name = BODY_FONT; r.font.size = Pt(8); r.font.color.rgb = GRAY_500


def tile_grid(doc, tiles, cols=2, title_size=9.5):
    rows = (len(tiles) + cols - 1) // cols
    t = doc.add_table(rows=rows, cols=cols)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_borders(t, ['top', 'left', 'bottom', 'right', 'insideH', 'insideV'], size=4, color='E2E8F0')
    for idx, (title, text) in enumerate(tiles):
        cell = t.cell(idx // cols, idx % cols)
        cell.width = Mm(86)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        r = p.add_run(title)
        r.bold = True; r.font.name = BODY_FONT; r.font.size = Pt(title_size); r.font.color.rgb = MIDNIGHT
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(3)
        r = p2.add_run(text)
        r.font.name = BODY_FONT; r.font.size = Pt(8.5); r.font.color.rgb = SLATE
    return t


def layer_card(doc, code, title, text, chips, badge_fill='050A14'):
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    set_borders(t, ['top', 'left', 'bottom', 'right'], size=4, color='E2E8F0')
    badge, main = t.cell(0, 0), t.cell(0, 1)
    badge.width = Mm(14)
    main.width = Mm(158)
    shade(badge, badge_fill)
    cell_text(badge, [(code, dict(bold=True, size=10, color=WHITE))], align=WD_ALIGN_PARAGRAPH.CENTER)
    p = main.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    r = p.add_run(title)
    r.bold = True; r.font.name = BODY_FONT; r.font.size = Pt(10.5); r.font.color.rgb = MIDNIGHT
    p2 = main.add_paragraph()
    add_runs(p2, text, size=9)
    p3 = main.add_paragraph()
    p3.paragraph_format.space_after = Pt(2)
    r = p3.add_run('  ·  '.join(chips))
    r.font.name = BODY_FONT; r.font.size = Pt(8); r.font.color.rgb = GRAY_500
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)


def main():
    doc = Document()
    doc.core_properties.title = 'PB Tech: Security Services Overview'
    doc.core_properties.author = 'ASG Security'
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(19)
    sec.top_margin, sec.bottom_margin = Mm(20), Mm(18)
    style = doc.styles['Normal']
    style.font.name = BODY_FONT
    style.font.size = Pt(10)
    style.font.color.rgb = SLATE

    cover(doc)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    stat_strip(doc)

    # 1 — The challenge
    overline(doc, 'The challenge')
    heading(doc, 'Technology retail faces unique risks')
    body(doc, 'Laptops, phones, components and accessories are compact, high-value and quick to resell. '
              'That stock profile shapes four risks every technology retailer manages daily:')
    tile_grid(doc, [
        ('Organised retail crime', 'Repeat offenders working across branches, with the same group appearing '
         'at different stores in the same week.'),
        ('Aggression at counters', 'Staff face anti-social behaviour, verbal abuse and physical threats, '
         'raising duty-of-care obligations as well as morale costs.'),
        ('After-hours attack', 'Stores and warehouses holding resaleable stock carry a standing burglary '
         'and ram-raid exposure outside trading hours.'),
        ('Fraud channels', 'Refund and warranty fraud, click-and-collect abuse and internal loss patterns '
         'that pure door security never sees.'),
    ])
    body(doc, '<b>Technology is a high-risk category.</b> Compact, serialised, high-resale stock is '
              'consistently among the most targeted categories across our own covert portfolio. The '
              'four-week pilot in the companion proposal puts the economics to the test on your own stores.')

    # 2 — Service lines
    overline(doc, 'Our service lines')
    heading(doc, 'Five capabilities, one provider')
    body(doc, 'Our integrated model combines covert operatives, uniformed officers, electronic systems, '
              'monitoring and a data layer under one provider: a security net that catches what cameras '
              'cannot, and prevents what a uniform alone never will. The companion proposal configures '
              'these capabilities specifically for PB Tech.')
    layer_card(doc, 'LP', 'Covert Loss Prevention',
               'Plain-clothes operatives blend with customers, walking the floor in minimum two-person '
               'teams. They identify, monitor and engage offenders where it is safe and lawful, recovering '
               'stock before it leaves the store.',
               ['Floor-walking', '2-person teams', 'Crime-intelligence trained', 'Real-time reporting'])
    layer_card(doc, 'UG', 'Uniformed Guard Presence',
               'Visible, professional officers at key entry and exit points. They deter opportunistic '
               'theft, assist staff, manage difficult situations and give customers and staff an immediate '
               'security presence.',
               ['Visible deterrent', 'Entry control', 'Staff support', 'De-escalation'])
    layer_card(doc, 'ES', 'Electronic Security Systems',
               'EAS theft-detection gates, product protection tagging, display alarm systems and anti-theft '
               'retail infrastructure, designed, installed and maintained as part of one integrated '
               'programme.',
               ['EAS gates', 'Product tagging', 'Display alarms', 'CCTV'])
    layer_card(doc, 'MN', 'Monitoring and Alarm Response',
               'Around-the-clock monitoring of alarm activations, with AI video verification where live '
               'footage is available and an agreed response protocol behind it, backed by patrol or '
               'keyholder attendance.',
               ['24/7 monitoring', 'Video verification', 'Alarm response', 'Camera health'])
    layer_card(doc, 'DL', 'Data-Led Coverage',
               'Coverage weighted to measured loss and reviewed weekly, so hours sit at the stores where '
               'offending is actually happening rather than where a tender fixed them.',
               ['Weekly re-weighting', 'Severity grading', 'Cost vs budget'])
    body(doc, '<b>Specialist support.</b> Two further capabilities round out the service: '
              '<b>Investigations and Intelligence</b> (internal investigations, serialised-stock tracing, '
              'refund and warranty fraud, click-and-collect fraud screening) and <b>Surge and Event '
              'Support</b> (launch-day crowd management, promotional peaks, multi-store simultaneous '
              'activation across regions).')
    body(doc, '<b>Optional roadmap: facial recognition.</b> Watchlist alerting for serious repeat '
              'offenders is available as a separately assessed option only: a documented necessity and '
              'proportionality assessment and privacy assessment come first, PB Tech owns the go or no-go '
              'decision, every alert is human-verified before any response, and strict enrolment, '
              'retention and signage controls apply under the Biometric Processing Privacy Code 2025.')

    # 3 — How we operate
    overline(doc, 'Operational model')
    heading(doc, 'How we operate')
    tile_grid(doc, [
        ('01  Minimum two-person teams', 'Every covert shift uses at least two operatives in tandem: one '
         'monitors, one positions for the approach. Safe, lawful engagements with a corroborating witness '
         'at every event.'),
        ('02  Flexible shift structure', 'Shifts are structured to keep operatives sharp across the full '
         'trading day.'),
        ('03  Floor walkers, not screens', 'Operatives walk the floor among your customers, observing '
         'behaviour, watching concealment points and positioning near high-risk product.'),
        ('04  Dedicated communications', 'Covert and uniformed teams stay in continuous contact for '
         'real-time coordination and rapid response when an incident unfolds.'),
        ('05  Agreed SOPs before day one', 'Engagement protocols, escalation paths and reporting lines are '
         'agreed upfront. Engagements occur only where safe and lawful: our staff claim no detention '
         'powers beyond any citizen’s, and trespass action is taken only as PB Tech’s authorised agent.'),
        ('06  Weekly roster, verified hours', 'A detailed roster is issued weekly and every hour billed is '
         'timeclock-verified. Full transparency on who is working, where and when.'),
    ])

    # 4 — Electronic systems
    overline(doc, 'Electronic security')
    heading(doc, 'Integrated electronic systems')
    body(doc, 'Beyond physical guarding, we design, install and maintain electronic security '
              'infrastructure tailored to your store layout and loss profile. One provider managing both '
              'physical and electronic security means seamless integration and single-point accountability.')
    bullet(doc, '<b>EAS theft-detection gates (RF and AM):</b> entry and exit detection calibrated for '
                'high-traffic retail, triggering guard response on activation.')
    bullet(doc, '<b>Product protection tagging:</b> hard tags, cable locks and spider wraps for laptops, '
                'phones and high-risk lines. Targeted protection where losses are highest.')
    bullet(doc, '<b>Display alarm systems:</b> alarmed display protection for demonstration devices that '
                'keeps product presentation open to customers.')
    bullet(doc, '<b>Anti-theft infrastructure:</b> bollards, secure fixtures and access control for '
                'high-value stock areas, service-centre custody areas and back-of-house.')
    bullet(doc, '<b>CCTV:</b> we integrate with your existing camera estate first and specify new hardware '
                'only where coverage gaps create risk; supply, installation and servicing are in-house. '
                'The Mosquito anti-loitering system is available where exterior congregation is a problem.')

    # 5 — Reporting
    overline(doc, 'Technology')
    heading(doc, 'Intelligence-led reporting')
    tile_grid(doc, [
        ('Deployment & Recovery Tracker', 'Hours, recoveries, incidents, offenders, deterrents and '
         'trespasses per store, with cost tracked against budget.'),
        ('National summary', 'Store, region and national totals with week-on-week movements explained, '
         'not just reported.'),
        ('Safety report', 'Severity spread on our 1 to 5 scale, with serious incidents narrated and '
         'Police referrals tracked.'),
        ('Verified attendance', 'Every hour billed reconciles to timeclock records. No guesswork, no '
         'ghost shifts.'),
    ])
    body(doc, '<b>Crime-intelligence integration.</b> Operatives are trained on retail crime-intelligence '
              'platforms such as Auror. Where authorised by you and consistent with your privacy policy, '
              'incidents are recorded in agreed systems and structured evidence packs prepared for Police '
              'consideration, subject always to Police evidential requirements and discretion.')
    body(doc, '<b>Reporting cycle.</b> Real-time incident logging, a weekly pack every Tuesday, a monthly '
              'review of trends and recommendations, and a quarterly strategy session where coverage is '
              're-tiered from the data.')

    # 6 — Compliance
    overline(doc, 'Compliance')
    heading(doc, 'Licensed, trained and compliant')
    tile_grid(doc, [
        ('Personnel standards', '✓ Required current licence or Certificate of Approval under the '
         'Private Security Personnel and Private Investigators Act 2010, in the classes each role '
         'requires\n✓ Conflict resolution and de-escalation training\n✓ First aid\n'
         '✓ Site-specific SOP training before any shift\n✓ Pre-employment vetting including '
         'Police checks'),
        ('Operational quality', '✓ Timeclock-verified attendance at every shift\n✓ Digital, '
         'time-stamped incident logging\n✓ Weekly performance reporting\n✓ Monthly trend '
         'analysis and KPIs\n✓ Regular rotation to prevent complacency'),
        ('Insurance and liability', '✓ Public liability cover held by both entities\n✓ Statutory '
         'liability cover held by both entities\n✓ Certificates of currency and limits available on '
         'request'),
        ('Health and safety', '✓ We meet our PCBU duties under the Health and Safety at Work Act and '
         'consult, cooperate and coordinate with PB Tech on overlapping duties\n✓ Incident and '
         'near-miss reporting\n✓ De-escalation first; physical contact a last resort\n✓ 24/7 '
         'operations support line'),
    ])

    # 7 — Results
    overline(doc, 'Proven results')
    heading(doc, 'Delivering results across New Zealand')
    body(doc, 'We protect major national retail networks with covert and uniformed services across '
              'grocery, hardware and general merchandise, alongside precinct, education, hospitality, '
              'property and corporate contracts.')
    bullet(doc, '<b>Covert loss prevention:</b> more than $1 million in verified recoveries since July '
                '2025, reported to clients weekly, across 100+ retail sites.')
    bullet(doc, '<b>Precinct and urban security:</b> a flagship multi-year city-centre precinct contract '
                'has recorded zero use of force across every aggression incident in its life, with '
                'thousands of patrol routes and assistance calls completed.')
    bullet(doc, '<b>Integrated electronic security:</b> EAS, product protection, CCTV and access control '
                'designed, installed and maintained by the same group that provides the officers.')

    # 8 — People
    overline(doc, 'Our people')
    heading(doc, 'Leadership team')
    if TEAM_PHOTO.exists():
        doc.add_picture(str(TEAM_PHOTO), width=Mm(120))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    body(doc, 'Our leadership experience is anchored by SLS Retail Theft Control’s continuous retail '
              'loss-prevention operation since 2006. The group is led by <b>Ahmad Duais</b> (Chief '
              'Executive), with <b>Mark Clansey</b> (General Manager) and <b>Hasi Nabawi</b> (Operations '
              'Manager) running national operations.')
    bullet(doc, '<b>Dedicated account manager</b> named at engagement: a single point of contact.')
    bullet(doc, '<b>Regional operations coverage</b> from our Auckland, Waikato, Wellington and '
                'Christchurch bases.')
    bullet(doc, '<b>24/7 operations support line</b> for urgent escalations.')
    bullet(doc, '<b>In-house training academy;</b> turnover across our permanent security workforce ran '
                'under 5% in the last 12 months.')
    bullet(doc, '<b>No subcontracting outside the group:</b> every officer is directly employed by an ASG '
                'Group company.')

    # 9 — Engagement
    overline(doc, 'Getting started')
    heading(doc, 'How an engagement starts')
    bullet(doc, '<b>1. Audit:</b> site walkthroughs and a loss-data workshop under NDA, producing a '
                'store-by-store coverage plan and risk tiering.')
    bullet(doc, '<b>2. SOPs:</b> engagement protocols, escalation paths and reporting lines agreed before '
                'day one.')
    bullet(doc, '<b>3. Pilot:</b> a four-week covert pilot in your highest-loss stores, reported from the '
                'first shift.')
    bullet(doc, '<b>4. Scale:</b> review on four weeks of real data, then staged rollout nationally.')
    body(doc, 'The full offer, guarantee, implementation phases and timing are set out in the companion '
              'proposal document.')

    # Contact block
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    cell = t.cell(0, 0)
    cell.width = Mm(172)
    shade(cell, '050A14')
    set_borders(t, ['top', 'left', 'bottom', 'right'], size=4, color='050A14')
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run('Ready to protect your stores.')
    r.bold = True; r.font.name = BODY_FONT; r.font.size = Pt(13); r.font.color.rgb = WHITE
    p2 = cell.add_paragraph()
    r = p2.add_run('Ahmad Duais  ·  Chief Executive, Absolute Security Group')
    r.font.name = BODY_FONT; r.font.size = Pt(10); r.font.color.rgb = GRAY_400
    p3 = cell.add_paragraph()
    r = p3.add_run('ahmad@absolutesecuritygroup.co.nz  ·  absolutesecuritygroup.co.nz  ·  slsrtc.co.nz')
    r.font.name = BODY_FONT; r.font.size = Pt(10); r.font.color.rgb = GRAY_400
    p4 = cell.add_paragraph()
    p4.paragraph_format.space_after = Pt(10)
    r = p4.add_run('Reply to this overview or call to arrange a site visit.')
    r.font.name = BODY_FONT; r.font.size = Pt(9); r.font.color.rgb = GRAY_500

    # Footer
    fp = sec.footer.paragraphs[0]
    fp.paragraph_format.tab_stops.add_tab_stop(Mm(170), WD_TAB_ALIGNMENT.RIGHT)
    r = fp.add_run('ASG Security  ·  Commercial in confidence\t')
    r.font.name = BODY_FONT; r.font.size = Pt(7.5); r.font.color.rgb = GRAY_500
    rp = fp.add_run('Page ')
    rp.font.name = BODY_FONT; rp.font.size = Pt(7.5); rp.font.color.rgb = GRAY_500
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    fp._p.append(fld)

    doc.save(str(OUT))
    print(f'OK {OUT}')


if __name__ == '__main__':
    main()
