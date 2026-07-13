#!/usr/bin/env python3
"""
1st Call Recruitment — earn-out proposal (2.69x EBITDA structure).
Builds:
  1) Stacked schedule chart ($11.5m-$13.5m): settlement + earn-out -> total  (PNG + PDF)
  2) Proposal email to Greg Dunn with the chart embedded                      (DOCX)

Structure: Settlement $750k fixed + 12-month earn-out, total = 2.69x maintainable
EBITDA, floored at $750k, capped at $1.6m. EBITDA = 0.2005*Rev - 2,093,288.

Deliverables -> ~/Desktop/Acquisitions/1st-Call-Recruitment/Working Files/
"""
import os
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

OUT = os.path.expanduser("~/Desktop/Acquisitions/1st-Call-Recruitment/Working Files")
PNG = os.path.join(OUT, "1CR_Earnout_Schedule_Chart.png")
PDF = os.path.join(OUT, "1CR_Earnout_Schedule_Chart.pdf")
DOCX = os.path.join(OUT, "1CR_Earnout_Proposal_Email_Greg.docx")

NAVY = "#13314F"; TEAL = "#1C8C8C"; GREEN = "#2E8B57"; GREY = "#9AA6B2"
RED = "#C0392B"; AMBER = "#B5762B"; INK = "#1A1A1A"

plt.rcParams.update({
    "font.family": "DejaVu Sans", "axes.edgecolor": "#D5DCE3",
    "text.color": INK, "axes.labelcolor": INK,
    "xtick.color": "#5A6B7B", "ytick.color": "#5A6B7B",
})

GP, OPEX, SETT, MULT, CAP = 0.2005, 2_093_288, 750_000, 2.69, 1_600_000
def eb(R):  return GP * R - OPEX
def total(R): return min(max(MULT * eb(R), SETT), CAP)
def earn(R):  return total(R) - SETT

pts = [
    (11_500_000, "$11.5m"), (12_000_000, "$12.0m"),
    (12_100_000, "$12.1m\nrun-rate"), (12_640_040, "$12.64m\nFY26"),
    (13_000_000, "$13.0m"), (13_200_000, "$13.2m"),
    (13_400_000, "$13.4m\nrecovery"), (13_500_000, "$13.5m"),
]

fig, ax = plt.subplots(figsize=(11.4, 7.2))
fig.subplots_adjust(left=0.085, right=0.965, top=0.8, bottom=0.13)

xs = range(len(pts))
sett_vals = [SETT / 1000] * len(pts)
earn_vals = [earn(R) / 1000 for R, _ in pts]

b1 = ax.bar(xs, sett_vals, width=0.6, color=NAVY, zorder=3, label="Settlement (paid at completion)")
b2 = ax.bar(xs, earn_vals, bottom=sett_vals, width=0.6, color=GREEN, alpha=0.85,
            zorder=3, label="12-month earn-out (revenue-linked)")

# cap line
ax.axhline(CAP / 1000, color=RED, lw=1.4, ls=(0, (5, 3)), zorder=4)
ax.text(0.0, CAP / 1000 + 16, r"Cap  \$1.6m", color=RED, fontsize=9.5,
        fontweight="bold", ha="left", va="bottom")

for i, (R, _) in enumerate(pts):
    tot = total(R) / 1000
    eo = earn(R) / 1000
    # total label on top
    fy = (abs(R - 12_640_040) < 1)
    ax.text(i, tot + 22, rf"\${tot:,.0f}k", ha="center", va="bottom",
            fontsize=9.0, fontweight="bold", color=NAVY, zorder=6)
    # earn-out value centred inside the green segment (skip if too small to hold text)
    if eo >= 60:
        ax.text(i, SETT/1000 + eo/2, rf"+\${eo:,.0f}k", ha="center", va="center",
                fontsize=(8.3 if eo >= 150 else 7.4), fontweight="bold",
                color="white", zorder=6)
    # FY26 highlight outline
    if fy:
        ax.bar(i, tot, width=0.64, fill=False, edgecolor=AMBER, lw=1.8, zorder=5)

# settlement label inside first bar
ax.text(0, SETT/1000/2, r"\$750k", ha="center", va="center", fontsize=8.5,
        fontweight="bold", color="white", zorder=6)

ax.set_xticks(list(xs))
ax.set_xticklabels([lab for _, lab in pts], fontsize=8.7)
ax.set_ylim(0, 1850)
ax.set_ylabel(r"Total consideration (\$000s)", fontsize=10.5)
ax.set_xlabel("Trailing-12-month revenue at the earn-out measurement", fontsize=10.5)
ax.yaxis.set_major_formatter(plt.FuncFormatter(lambda y, _: rf"\${y:,.0f}k"))
ax.grid(axis="y", color="#EEF2F5", lw=0.9, zorder=0)
for s in ("top", "right"):
    ax.spines[s].set_visible(False)
ax.legend(loc="upper left", frameon=False, fontsize=9.2, bbox_to_anchor=(0.0, 1.02))

fig.suptitle("1st Call Recruitment — Settlement + Revenue-Linked Earn-out",
             x=0.085, y=0.955, ha="left", fontsize=16, fontweight="bold", color=NAVY)
fig.text(0.085, 0.875,
         r"\$750k at completion, plus a 12-month earn-out that grows with the revenue the business delivers in its first year —"
         "\n"
         r"reaching the full \$1.6m only if revenue recovers to ~\$13.4m. Floored at the \$750k settlement; capped at \$1.6m.",
         ha="left", fontsize=9.6, color="#5A6B7B", linespacing=1.4)
fig.text(0.085, 0.028,
         r"Total consideration = settlement + 12-month earn-out, measured on trailing-12-month revenue after completion; "
         r"floored at the \$750k settlement and capped at \$1.6m.",
         ha="left", fontsize=7.5, color="#8895A2")

fig.savefig(PNG, dpi=300, facecolor="white"); print("Wrote", PNG)
fig.savefig(PDF, facecolor="white"); print("Wrote", PDF)

# ---- email -------------------------------------------------------------
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
st.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
NAVY_T = (0x13, 0x31, 0x4F)

def p(text="", bold=False, size=11, color=None, sa=6):
    par = doc.add_paragraph(); par.paragraph_format.space_after = Pt(sa)
    r = par.add_run(text); r.bold = bold; r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor(*color)
    return par

def bullet(text, lead=None):
    par = doc.add_paragraph(style="List Bullet"); par.paragraph_format.space_after = Pt(3)
    if lead:
        rr = par.add_run(lead); rr.bold = True
    par.add_run(text)
    return par

hdr = doc.add_paragraph(); hdr.paragraph_format.space_after = Pt(2)
for lbl, val in [("To:  ", "Greg Dunn — gregd@abcbusiness.co.nz"),
                 ("Cc:  ", "Tony Begbie — tonyb@abcbusiness.co.nz")]:
    rr = hdr.add_run(lbl); rr.bold = True; hdr.add_run(val + "\n")
subj = doc.add_paragraph(); subj.paragraph_format.space_after = Pt(10)
sr = subj.add_run("Subject:  1st Call Recruitment — a structure that pays Phill his $1.6m on recovery")
sr.bold = True; sr.font.color.rgb = RGBColor(*NAVY_T)

p("Hi Greg,")
p("Here’s a structure that closes the gap between our numbers and Phill’s $1.6m — by paying "
  "it in full if the recovery he believes in actually lands, funded out of that recovery "
  "rather than up front.")

p("The proposal", bold=True, size=12, color=NAVY_T, sa=4)
bullet("paid at completion, and Phill’s regardless of what follows.", lead="$750,000 at settlement:  ")
bullet("a single payment made at the first anniversary, calculated on the trailing-12-month revenue the business delivers in its first year under our ownership.", lead="12-month earn-out:  ")
bullet("the earn-out grows with that revenue — the stronger the recovery, the more it pays. Reach around $13.4m and it pays the full $1.6m.", lead="How it scales:  ")
bullet("the total is capped at $1.6m, and the settlement is the floor — the earn-out is never negative.", lead="Capped and floored:  ")

doc.add_picture(PNG, width=Inches(6.3))
cap = doc.paragraphs[-1]; cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

p("Why this should work for Phill", bold=True, size=12, color=NAVY_T, sa=4)
p("Phill’s case for $1.6m rests on the business recovering. This pays him exactly that — in "
  "full — if it does. He takes $750,000 in cash on day one and collects the balance twelve "
  "months later, out of the very revenue the price is based on. If the recovery is as real "
  "as he believes, he is no worse off than a $1.6m sale today; he simply receives part of it "
  "a year on, once it’s proven.")
p("From our side it lets us commit, because we’re paying for performance we can see rather "
  "than paying up for a forecast we can’t yet verify — we’re not looking to catch a falling "
  "knife. It puts the risk where the conviction is.")

p("What I’d need to finalise it", bold=True, size=12, color=NAVY_T, sa=4)
bullet("The FY25 monthly P&L (or at least April–May 2025), so we set the baseline like-for-like against the seasonal pattern rather than the peak.", lead="Baseline data:  ")
bullet("Agreement that the earn-out is measured on the transferred client book over the 12 months post-completion, and on how new business won in that period is treated.", lead="Mechanics:  ")
p("Happy to walk Phill or KPMG through the structure and the earnings work behind it on a "
  "call. I think this is the most honest way to bridge the gap — it gives Phill his number if "
  "the business delivers, and gives us a deal we can stand behind either way.")

p("Kind regards,", sa=2)
p("Ahmad Duais", bold=True, sa=0)
p("Absolute Security Group", sa=0, size=10, color=(0x5A, 0x6B, 0x7B))

doc.save(DOCX)
print("Wrote", DOCX)

# remove the superseded linear-ladder chart to avoid confusion
for f in ("1CR_Earnout_Ladder_Chart.png", "1CR_Earnout_Ladder_Chart.pdf"):
    fp = os.path.join(OUT, f)
    if os.path.exists(fp):
        os.remove(fp); print("Removed superseded", f)
