#!/usr/bin/env python3
"""v4 — strip IT in FULL (not needed post-acquisition): add back the entire Wages-IT
$116,614 (FY26) / $248,135 (FY25), keeping $0 forward IT. Raises FY26 normalised EBITDA to
~$401K (incl KiwiSaver) / ~$466K (ex-KiwiSaver). Idempotent over v3."""
import openpyxl
SRC = "/Users/ahmadduais/Desktop/Acquisitions/1st-Call-Recruitment/1CR_Acquisition_Model.xlsx"
MIRROR = "/Users/ahmadduais/Desktop/Acquisitions/1st-Call-Recruitment/FINAL Acquisition Pack/1CR_Acquisition_Model.xlsx"
from openpyxl.utils.cell import coordinate_to_tuple
wb = openpyxl.load_workbook(SRC)
def sset(ws, coord, val):
    tr, tc = coordinate_to_tuple(coord)
    for mr in list(ws.merged_cells.ranges):
        if mr.min_row <= tr <= mr.max_row and mr.min_col <= tc <= mr.max_col:
            ws.unmerge_cells(str(mr)); break
    ws[coord] = val
REV26, REV25, REV27 = 12640040, 17890076, 12100000
ws = wb["Normalisation & FY25"]
ws["C10"], ws["D10"] = 248135, 116614   # full IT strip (GL Wages-IT)
ws["F10"] = "Wages-IT stripped IN FULL — not required post-acquisition (ASG IT covers it). FY26 $116,614 / FY25 $248,135."
def col_sum(letter):
    return sum((ws[f"{letter}{r}"].value or 0) for r in range(8, 20))
fy25, fy26 = col_sum("C"), col_sum("D"); fy26_exks = fy26 + 65000
ws["C20"], ws["D20"], ws["E20"] = fy25, fy26, 410000
ws["F20"] = f"FY26 ~${fy26/1000:.0f}K (~${fy26_exks/1000:.0f}K ex-KiwiSaver). IT stripped in full. NOT $783K."
sset(ws, "C21", round(fy25/REV25,4)); sset(ws, "D21", round(fy26/REV26,4)); sset(ws, "E21", round(410000/REV27,4))

ws = wb["Mgmt Accounts Reconciliation"]
ws["C11"], ws["D11"] = "~$401K (full strip, IT fully removed)", "~$401–466K"
ws["F11"] = "Three accrual releases stripped + IT removed IN FULL ($116,614) + rent/FBT normalised − KiwiSaver = ~$401K (~$466K ex-KiwiSaver), NOT $783K."

ws = wb["Group Consolidated"]
ws["F9"], ws["G9"] = fy26, round(fy26/REV26,4)
ws["F10"] = 1436680 + fy26; ws["G10"] = round((1436680+fy26)/20191220,4)
ws["F18"], ws["G18"] = 410000, round(410000/REV27,4)
ws["F19"] = 1609082 + 410000; ws["G19"] = round((1609082+410000)/20557322,4)
ws["C24"] = f"$1.44M → ${(1436680+fy26)/1e6:.2f}M"; ws["F24"] = f"+{round(fy26/1436680*100)}%"
ws["C26"] = f"19.0% → {round((1436680+fy26)/20191220*100,1)}%"

ws = wb["Summary"]
ws["C12"], ws["D12"] = fy26, 410000
ws["C13"], ws["D13"] = fy26_exks, 470000
ws["C14"], ws["D14"] = round(fy26/REV26,4), round(410000/REV27,4)
ws["B16"] = f"RETURNS (corrected EBITDA ~${fy26/1000:.0f}K)"
ws["C17"] = f"{1100000/fy26:.2f}x @ $1.1M · {600000/fy26:.2f}x @ $600K target · {1600000/fy26:.2f}x @ his $1.6M"
ws["C19"] = fy25
wb.save(SRC); wb.save(MIRROR)
print(f"FY25={fy25:,.0f}  FY26={fy26:,.0f} (ex-KiwiSaver {fy26_exks:,.0f})")
print(f"Multiple: {1100000/fy26:.2f}x @$1.1M | {600000/fy26:.2f}x @$600K | {1600000/fy26:.2f}x @$1.6M")
