#!/usr/bin/env python3
"""v2 — correct the v1 bridge error. Strip the FULL actual P&L accrual credits on the
FY26 (declining) basis: Holiday Pay Accrual release -$154,914 and Temping Pays Accrual
reversal -$40,277 (do NOT net against FY25). KiwiSaver under-accrual -$65,000 shown as a
DD-contingent line. Corrected FY26 normalised EBITDA = $430,389 (~$495K ex-KiwiSaver),
GP ~20.0%. Sets absolute values (idempotent) over the v1-modified workbook."""
import openpyxl
from openpyxl.utils.cell import coordinate_to_tuple

SRC = "/Users/ahmadduais/Desktop/Acquisitions/1st-Call-Recruitment/1CR_Acquisition_Model.xlsx"
MIRROR = "/Users/ahmadduais/Desktop/Acquisitions/1st-Call-Recruitment/FINAL Acquisition Pack/1CR_Acquisition_Model.xlsx"
wb = openpyxl.load_workbook(SRC)

def sset(ws, coord, val):
    tr, tc = coordinate_to_tuple(coord)
    for mr in list(ws.merged_cells.ranges):
        if mr.min_row <= tr <= mr.max_row and mr.min_col <= tc <= mr.max_col:
            ws.unmerge_cells(str(mr)); break
    ws[coord] = val

REV26, REV25, REV27 = 12640040, 17890076, 12100000

# ---------- Normalisation & FY25 ----------
ws = wb["Normalisation & FY25"]
# GP$ / GP% — FY26 fully stripped (holiday+temping) = 20.05%; FY25 holiday-strip = 20.34%; FY27 flat 20.05%
GP26 = round(REV26*0.2005); GP25 = round(REV25*0.2034); GP27 = round(REV27*0.2005)
ws["C6"], ws["D6"], ws["E6"] = GP25, GP26, GP27
ws["C7"], ws["D7"], ws["E7"] = 0.2034, 0.2005, 0.2005
ws["F7"] = "FY26 ~20.0% (full strip of holiday-pay release -$154,914 AND temping-pays reversal -$40,277). Declining business → underwrite FY26, not the FY25 average."
ws["F6"] = "~20.0% maintainable (FY26 basis) — reported 21.6% is two accrual credits inside COGS, not margin"

# Bridge lines (FY26 = D, FY25 = C). Rows after v1: 16/17/18 = strip lines, 19 = EBITDA, 20 = margin
ws["B16"] = "− Holiday Pay Accrual release (strip balance-sheet credit from COGS)"
ws["C16"], ws["D16"] = -268156, -154914
ws["F16"] = "FY26 -$154,914 / FY25 -$268,156 — provision RELEASE as headcount contracts ~29%; non-cash, flatters GP"
ws["B17"] = "− Temping Pays Accrual reversal (FY26 credit; FY26 basis, NOT netted vs FY25)"
ws["C17"], ws["D17"] = 0, -40277
ws["F17"] = "FY26 -$40,277 credit flatters GP; declining-business basis → strip in full, don't average with FY25's +$40,277 charge"
ws["B18"] = "− KiwiSaver under-accrual (DD-contingent; hold at FY25 effective 1.89%)"
ws["C18"], ws["D18"] = 0, -65000
ws["F18"] = "FY26 effective 1.18% vs FY25 1.89% — apparent under-accrual; ~$65K returns if confirmed in DD"

def col_sum(letter):
    return sum((ws[f"{letter}{r}"].value or 0) for r in range(8, 19))
fy25, fy26 = col_sum("C"), col_sum("D")
fy26_exks = fy26 + 65000
ws["C19"], ws["D19"] = fy25, fy26
ws["E19"] = 400000  # FY27 flat-to-declining on the lower basis
ws["F19"] = f"FY26 ~${fy26/1000:.0f}K (~${fy26_exks/1000:.0f}K ex-KiwiSaver). FY27 flat-declining ~$400K. NOT $783K."
ws["C20"] = round(fy25/REV25, 4); ws["D20"] = round(fy26/REV26, 4); ws["E20"] = round(400000/REV27, 4)

sset(ws, "C23", f"• FY25 normalised EBITDA ~${fy25/1e6:.2f}M at $17.9M revenue = earnings power AT SCALE — NOT the basis. Revenue has since fallen 29%.")
sset(ws, "B24", f"• FY26 buyer-normalised EBITDA ~${fy26/1000:.0f}K (≈${fy26_exks/1000:.0f}K ex-KiwiSaver) after stripping the FULL holiday-pay release (-$154,914) and temping-pays reversal (-$40,277) from COGS and correcting the IT & rent add-backs. GP ~20.0%, not 21%+. The earlier $783K embedded the accrual credits and over-stated add-backs.")
sset(ws, "B25", "• FY27 modelled FLAT-to-DECLINING ~$400K. FY26 fell -29.3% YoY; trailing-12mo ~$12.1M is ~4% below FY26; no contracted pipeline. We underwrite the declining FY26 base, not a recovery.")

# ---------- Mgmt Accounts Reconciliation ----------
ws = wb["Mgmt Accounts Reconciliation"]
ws["C6"], ws["E6"] = "~20.0% (maintainable, FY26 basis)", "⚠ FLATTERED"
ws["F6"] = "Reported GP is flattered by TWO accrual credits inside COGS — Holiday Pay Accrual release (FY26 -$154,914) and Temping Pays Accrual reversal (FY26 -$40,277). Strip both on the FY26 declining basis → maintainable GP ~20.0%."
ws["C11"], ws["D11"], ws["E11"] = "~$430K (fully accrual-stripped, FY26 basis)", "~$400–450K trailing", "⚠ CORRECTED DOWN"
ws["F11"] = "Full strip of holiday-pay + temping-pays credits + corrected IT/rent + KiwiSaver = ~$430K (~$495K ex-KiwiSaver), not $783K. Underwrite FY26 (declining), not the FY25 average."
ws["B14"] = ("Net: declining business — underwrite FY26. GP is ~20.0% once BOTH accrual credits (holiday-pay release -$154,914, temping-pays reversal -$40,277) are stripped from COGS; the reported 21.6%/20.8% is not operating margin. Buyer-normalised FY26 EBITDA ~$430K (~$495K ex-KiwiSaver), NOT $783K. '~$148K interest upside' is a phantom and is deleted.")

# ---------- Group Consolidated ----------
ws = wb["Group Consolidated"]
ws["D9"], ws["E9"], ws["F9"], ws["G9"] = GP26, 0.2005, fy26, round(fy26/REV26, 4)
ws["D10"] = 1621420 + GP26
ws["F10"] = 1436680 + fy26
ws["E10"] = round((1621420 + GP26)/20191220, 4)
ws["G10"] = round((1436680 + fy26)/20191220, 4)
ws["B11"] = f"1st Call at FY26 buyer-normalised ~${fy26/1000:.0f}K (full accrual strip, GP ~20.0%, declining revenue). Earlier $783K relied on accrual credits + over-stated add-backs."
ws["D18"], ws["E18"], ws["F18"], ws["G18"] = GP27, 0.2005, 400000, round(400000/REV27, 4)
ws["D19"] = 1815990 + GP27
ws["F19"] = 1609082 + 400000
ws["G19"] = round((1609082 + 400000)/20557322, 4)
ws["B20"] = "Existing group +12% (CSFpace); 1st Call FLAT-to-DECLINING (no recovery)."
ws["C24"] = f"$1.44M → ${(1436680+fy26)/1e6:.2f}M"
ws["F24"] = f"+{round(fy26/1436680*100)}%"
ws["C26"] = f"19.0% → {round((1436680+fy26)/20191220*100,1)}%"

# ---------- Summary ----------
ws = wb["Summary"]
ws["C5"] = "~$750K target (open $650K · walk-away $850K)"
ws["C6"] = "Offered $1.1M (REJECTED); vendor seeks $1.6M (≈3.7x corrected EBITDA — not supportable)"
ws["C11"], ws["D11"] = REV26, REV27
ws["C12"], ws["D12"] = fy26, 400000
ws["B12"] = "Normalised EBITDA (FY26 basis, full accrual strip)"
ws["C13"], ws["D13"] = fy26_exks, 465000
ws["B13"] = "If KiwiSaver DD-benign (ex under-accrual)"
ws["C14"], ws["D14"] = round(fy26/REV26, 4), round(400000/REV27, 4)
ws["B16"] = f"RETURNS (corrected EBITDA ~${fy26/1000:.0f}K)"
ws["C17"] = f"{1100000/fy26:.2f}x @ $1.1M · {850000/fy26:.2f}x @ $850K · {1600000/fy26:.2f}x @ his $1.6M"
ws["D17"] = ""
ws["C18"] = "n/a — at his $1.6M, after-tax payback >5yrs"
ws["C19"] = fy25

wb.save(SRC); wb.save(MIRROR)
print(f"FY25={fy25:,.0f}  FY26={fy26:,.0f} (ex-KiwiSaver {fy26_exks:,.0f})  FY27=400,000")
print(f"GP%: FY26 20.05% (${GP26:,.0f})  FY25 20.34% (${GP25:,.0f})")
print(f"Multiple: {1100000/fy26:.2f}x @$1.1M | {850000/fy26:.2f}x @$850K | {1600000/fy26:.2f}x @$1.6M")
