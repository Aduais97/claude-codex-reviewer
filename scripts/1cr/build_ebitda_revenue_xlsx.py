#!/usr/bin/env python3
"""
1st Call Recruitment — EBITDA-by-revenue table as a live Excel workbook.
Maintainable EBITDA(R) = GP% * R - fixed OpEx ; KiwiSaver-normalised = that - haircut.
All cells are formulas off an assumptions block so inputs flow through.

Deliverable -> ~/Desktop/Acquisitions/1st-Call-Recruitment/Working Files/
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.utils import get_column_letter

OUT = os.path.expanduser("~/Desktop/Acquisitions/1st-Call-Recruitment/Working Files")
XLSX = os.path.join(OUT, "1CR_EBITDA_by_Revenue.xlsx")

NAVY = "13314F"
TEAL = "1C8C8C"
LIGHT = "EEF3F6"
BAND = "F4F7F9"
INK = "1A1A1A"
GREY = "5A6B7B"

wb = Workbook()
ws = wb.active
ws.title = "EBITDA by Revenue"
ws.sheet_view.showGridLines = False

thin = Side(style="thin", color="D5DCE3")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

def cell(r, c, val=None, bold=False, size=11, color=INK, fill=None,
         align="left", fmt=None, border_on=False, italic=False):
    cl = ws.cell(row=r, column=c)
    if val is not None:
        cl.value = val
    cl.font = Font(name="Calibri", bold=bold, size=size, color=color, italic=italic)
    cl.alignment = Alignment(horizontal=align, vertical="center", wrap_text=False)
    if fill:
        cl.fill = PatternFill("solid", fgColor=fill)
    if fmt:
        cl.number_format = fmt
    if border_on:
        cl.border = border
    return cl

# ---- title -------------------------------------------------------------
cell(1, 1, "1st Call Recruitment (JCR 2006 Ltd)", bold=True, size=15, color=NAVY)
cell(2, 1, "Maintainable EBITDA by revenue  —  FY26 normalised basis", size=11, color=GREY)

# ---- assumptions block -------------------------------------------------
cell(4, 1, "Assumptions", bold=True, size=11, color=NAVY)
cell(5, 1, "Maintainable gross margin", color=INK)
cell(5, 2, 0.2005, align="right", fmt="0.00%")
cell(6, 1, "Fixed operating costs (below GP)", color=INK)
cell(6, 2, 2093288, align="right", fmt='#,##0')
cell(7, 1, "KiwiSaver-normalisation haircut", color=INK)
cell(7, 2, 65000, align="right", fmt='#,##0')
for r in (5, 6, 7):
    cell(r, 2, fill=LIGHT)
    ws.cell(row=r, column=2).border = border

GPCELL, OPEXCELL, KSCELL = "$B$5", "$B$6", "$B$7"

# ---- table header ------------------------------------------------------
H = 9
headers = ["Revenue", "EBITDA\n(maintainable)", "EBITDA\n(KiwiSaver-normalised)", "Note"]
for i, h in enumerate(headers, start=1):
    c = cell(H, i, h, bold=True, size=10.5, color="FFFFFF", fill=NAVY,
             align="center", border_on=True)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws.row_dimensions[H].height = 32

rows = [
    (12_000_000, ""),
    (12_100_000, "run-rate"),
    (12_400_000, ""),
    (12_600_000, ""),
    (12_640_040, "FY26 actual"),
    (12_800_000, ""),
    (13_000_000, ""),
    (13_200_000, ""),
    (13_400_000, "recovery scenario"),
]

r = H + 1
for rev, note in rows:
    band = BAND if note in ("FY26 actual",) else None
    cell(r, 1, rev, align="right", fmt='$#,##0', fill=band, border_on=True,
         bold=(note == "FY26 actual"))
    # EBITDA = GP% * Rev - fixed opex
    cell(r, 2, f"={GPCELL}*A{r}-{OPEXCELL}", align="right", fmt='$#,##0',
         fill=band, border_on=True, bold=(note == "FY26 actual"), color=NAVY)
    # KiwiSaver-normalised = above - haircut
    cell(r, 3, f"=B{r}-{KSCELL}", align="right", fmt='$#,##0',
         fill=band, border_on=True, color=TEAL)
    cell(r, 4, note, italic=True, color=GREY, fill=band, border_on=True)
    r += 1

# ---- footnotes ---------------------------------------------------------
fn = r + 1
cell(fn, 1,
     "Basis: FY26 normalised EBITDA $441k (owner salary, IT wages, FBT, disposal gain and the three accrual",
     size=8.5, color=GREY)
cell(fn + 1, 1,
     "releases adjusted; $50k sublease retained). Each $1.0m of revenue adds ~$200k of EBITDA (20% GP, fixed cost).",
     size=8.5, color=GREY)
cell(fn + 2, 1,
     "KiwiSaver-normalised line restates the FY26 1.18% effective rate toward FY25's 1.89% — a DD item, not settled.",
     size=8.5, color=GREY)
cell(fn + 3, 1,
     "Source: FY26 Xero Profit & Loss, JCR 2006 Limited, year ended 31 March 2026 (draft).",
     size=8.5, color=GREY)

# ---- widths ------------------------------------------------------------
ws.column_dimensions["A"].width = 20
ws.column_dimensions["B"].width = 20
ws.column_dimensions["C"].width = 26
ws.column_dimensions["D"].width = 20
ws.freeze_panes = "A10"

wb.save(XLSX)
print("Wrote", XLSX)
